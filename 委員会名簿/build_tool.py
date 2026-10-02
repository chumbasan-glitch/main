import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.formatting.rule import FormulaRule, Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils import get_column_letter as CL

OUT = '/home/user/main/委員会名簿/委員会名簿作成ツール（テストデータ入り）.xlsx'
FONT = '游ゴシック'
NC, NCLS, NR, NM = 20, 8, 45, 50      # 委員会数, 1学年の組数, 1組の人数, 1委員会の人数
random.seed(21)

f = lambda **k: Font(name=FONT, **{'size': 11, **k})
YEL = PatternFill('solid', fgColor='FFF2CC')
PUR = PatternFill('solid', fgColor='E4DFEC')
GRY = PatternFill('solid', fgColor='EDEDED')
HDR = PatternFill('solid', fgColor='DDEBF7')
RED = PatternFill('solid', fgColor='F8CBAD')
LRED = PatternFill('solid', fgColor='FCE4D6')
thin = Side(style='thin', color='808080')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal='center', vertical='center')
LFT = Alignment(horizontal='left', vertical='center')
SMALLF = dict(size=9, color='808080')

wb = Workbook()
def sheet(name, first=False):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name
    return ws
def put(ws, ref, v, font=None, fill=None, al=None, box=False, unlock=False):
    c = ws[ref]; c.value = v; c.font = font or f()
    if fill: c.fill = fill
    if al: c.alignment = al
    if box: c.border = BOX
    if unlock: c.protection = Protection(locked=False)
    return c
def protect(ws, filt=False):
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
def cf(ws, rng, formula, fill=None, font=None, border=None):
    if fill is not None:   # 条件付き書式のぬりつぶしは、Excelでは bgColor を見るので両方に入れる
        c = 'FF' + fill.fgColor.rgb[-6:]
        fill = PatternFill(fill_type='solid', fgColor=c, bgColor=c)
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=fill, font=font, border=border))
def page(ws, size, orient, tall=1):
    ws.page_setup.paperSize = size; ws.page_setup.orientation = orient
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = tall
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
ROLE = {'club': ('【入力：委員会担当】うす紫のセル', PUR, '7030A0'),
        'tan': ('【入力：担任】黄色のセル', YEL, 'FFC000'),
        'none': ('【入力なし】見る・印刷するだけ', GRY, 'A6A6A6')}
def role(ws, ref, kind):
    t, fill, tab = ROLE[kind]
    put(ws, ref, t, Font(name=FONT, size=10, bold=True), fill, LFT, True)
    ws.sheet_properties.tabColor = tab

S = '設定'
COLORS = [('黒', '000000'), ('赤', 'D00000'), ('青', '0050D0'), ('緑', '008A3E'),
          ('黄', 'BF8F00'), ('橙', 'E46C0A'), ('紫', '7030A0'), ('茶', '8B4513')]
COLOR_LIST = '"' + ','.join(n for n, _ in COLORS) + '"'
def colorcf(ws, rng, expr):
    for name, hexc in COLORS:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{expr}="{name}"'], font=Font(name=FONT, bold=True, color=hexc)))
CN = f'{S}!$I$11:$I$30'                     # 委員会名
CLR = lambda c: f'IFERROR(INDEX({S}!$D$11:$D$26,MATCH({c},{S}!$C$11:$C$26,0)),"")'   # 表示→色

# ---------------- テストデータ ----------------
COMM = [('代表', '6梅', '会議室'), ('集会', '6月', '体育館'), ('放送', '5竹', '放送室'), ('音楽', '5松', '音楽室'),
        ('飼育', '6竹', '飼育小屋'), ('図書', '5梅', '図書室'), ('体育', '6松', '体育倉庫前'), ('保健', '5月', '保健室'),
        ('環境', '6竹', '理科室'), ('給食', '5松', '家庭科室'), ('掲示', '6月', '図工室'), ('広報', '5竹', 'パソコン室')]
SIZE = {'代表': 8, '集会': 12, '放送': 12, '音楽': 8, '飼育': 16, '図書': 12, '体育': 13, '保健': 8, '環境': 10, '給食': 12, '掲示': 9, '広報': 10}
CLASSES = ['松', '竹', '梅', '月']
TESTCOLOR = {'松': '緑', '竹': '青', '梅': '赤', '月': '黄'}
SUR = '佐藤 鈴木 高橋 田中 伊藤 渡辺 山本 中村 小林 加藤 吉田 山田 佐々木 山口 松本 井上 木村 林 斎藤 清水 山崎 森 池田 橋本 阿部 石川 山下 中島 石井 小川 前田 岡田 長谷川 藤田 後藤 近藤 村上 遠藤 青木 坂本'.split()
GIV = '蓮 陽翔 湊 蒼 樹 大和 悠真 陽太 朝陽 律 颯太 結翔 優斗 奏太 新 陸 海斗 健太 拓海 航 陽葵 凛 結菜 芽依 葵 紬 澪 結衣 咲良 美月 心春 杏 莉子 彩花 花音 柚葉 千尋 楓 真央 さくら'.split()
used = set()
def newname():
    while True:
        n = random.choice(SUR) + '　' + random.choice(GIV)
        if n not in used: used.add(n); return n
DATA = {}                                   # (学年, 組) -> [(名前, 委員会)]
for g in (6, 5):
    pool = [n for n, k in SIZE.items() for _ in range(k)]
    random.shuffle(pool)
    pool = pool[:120] + [random.choice(list(SIZE)) for _ in range(max(0, 120 - len(pool)))]
    for i, c in enumerate(CLASSES):
        DATA[(g, c)] = [[newname(), pool[i * 30 + j]] for j in range(30)]
for (g, c, j) in [(6, '竹', 5), (6, '月', 18), (5, '松', 9), (5, '梅', 22)]:
    DATA[(g, c)][j - 1][1] = None           # 委員会が未入力
DATA[(5, '月')][11][1] = '保険'             # 設定にない委員会名（打ちまちがい）

# ---------------- 1. 使い方 ----------------
ws = sheet('使い方', first=True)
ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 14; ws.column_dimensions['C'].width = 100
put(ws, 'B1', '委員会名簿作成ツール　使い方', f(size=16, bold=True))
lines = [
 ('■ 色のきまり（入力する人で色が分かれています）', 'h'),
 ('@PUR', ''), ('@YEL', ''),
 ('それ以外のセルには数式が入っているので保護しています（「校閲」→「シート保護の解除」で外せます。パスワードなし）。各シートの左上にも、だれが入力するシートかを表示しています。', ''),
 ('', ''),
 ('■ 【委員会担当】新年度のはじめにすること', 'h'),
 ('① 【設定】年度、学年（6・5）、組（名前・色・人数）、委員会（名前・担当・活動場所）を入力します。', ''),
 ('　 ・組の表に書いた組の数だけ、6年・5年シートに組の枠ができます。「人数」の数だけ、番号の枠ができます。', ''),
 ('　 ・委員会を増やす：空いている行に書きこむ　　減らす：行の中身を消す　　名前を変える：書きかえる', ''),
 ('', ''),
 ('■ 【担任】クラスで委員会を決めたらすること', 'h'),
 ('② 【6年】【5年】自分の組の列に、子どもの「名前」と「委員会」（プルダウン）を入れます。番号は自動で入ります。', ''),
 ('　 ・名前は、名簿からコピーして「値の貼り付け」をしてもかまいません。', ''),
 ('　 ・名前があるのに委員会が空欄の所は、うすい赤になります。設定にない委員会名は赤になります。', ''),
 ('', ''),
 ('■ 【委員会担当】全部の組が入力し終わったらすること', 'h'),
 ('③ 【人数チェック】委員会ごとの人数（6年・5年）と、委員会が未入力の子の数を確かめます。', ''),
 ('④ 【一覧名簿】【委員会①〜⑳】を印刷します。6年→5年、組の順、番号順に並びます。', ''),
 ('　 ・一覧名簿はA3横。上段（①〜⑩）が1ページ目、下段（⑪〜⑳）が2ページ目です。委員会が10以下なら1ページ目だけ印刷します。', ''),
 ('　 ・委員会①〜⑳のタブ名は、右クリック→「名前の変更」で「飼育」などに変えてもかまいません。', ''),
 ('', ''),
 ('■ 上限', 'h'),
 ('委員会 20／1学年の組 8／1組 45人／1つの委員会 50人', ''),
 ('', ''),
 ('■ テストデータについて', 'h'),
 ('名前はすべて架空です。本番で使うときは、【6年】【5年】の名前と委員会を消してから使ってください。', ''),
]
for i, (t, k) in enumerate(lines, start=3):
    if t == '@PUR':
        put(ws, f'B{i}', 'うす紫', f(bold=True), PUR, CEN, True); put(ws, f'C{i}', '＝委員会担当が入力する所（設定）')
    elif t == '@YEL':
        put(ws, f'B{i}', '黄色', f(bold=True), YEL, CEN, True); put(ws, f'C{i}', '＝担任が入力する所（6年・5年シートの名前と委員会）')
    elif k == 'h':
        put(ws, f'B{i}', t, f(bold=True, size=12))
    else:
        put(ws, f'C{i}' if t.startswith('　') else f'B{i}', t)
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = 'A6A6A6'
page(ws, ws.PAPERSIZE_A4, 'landscape', 0)

# ---------------- 2. 設定 ----------------
ws = sheet(S)
for col, w in zip('ABCDEFGHIJK', [8, 9, 13, 7, 7, 3, 3, 6, 16, 12, 16]):
    ws.column_dimensions[col].width = w
put(ws, 'A1', '設定', f(size=14, bold=True))
ws.merge_cells('D1:K1'); role(ws, 'D1', 'club')
put(ws, 'A3', '年度', f(bold=True)); put(ws, 'B3', '令和8年度', fill=PUR, box=True, unlock=True); ws.merge_cells('B3:C3')
put(ws, 'A4', '呼び名', f(bold=True)); put(ws, 'B4', '委員会', fill=PUR, box=True, unlock=True)
put(ws, 'A6', '学年の設定', f(bold=True))
put(ws, 'A7', '上の学年', fill=GRY, al=CEN, box=True); put(ws, 'B7', 6, fill=PUR, al=CEN, box=True, unlock=True)
put(ws, 'A8', '下の学年', fill=GRY, al=CEN, box=True); put(ws, 'B8', 5, fill=PUR, al=CEN, box=True, unlock=True)
put(ws, 'A9', '組の設定（上から並び順。学年ごとに最大8組）', f(bold=True))
for c, h in zip('ABCDE', ['学年', '組', '名簿での表示', '色', '人数']):
    put(ws, f'{c}10', h, f(bold=True), HDR, CEN, True)
for i in range(16):
    r = 11 + i
    put(ws, f'A{r}', '=$B$7' if i < 8 else '=$B$8', fill=GRY, al=CEN, box=True)
    cls = CLASSES[i % 8] if i % 8 < 4 else None
    put(ws, f'B{r}', cls, fill=PUR, al=CEN, box=True, unlock=True)
    put(ws, f'C{r}', f'=IF(B{r}="","",A{r}&B{r})', fill=PUR, al=CEN, box=True, unlock=True)
    put(ws, f'D{r}', TESTCOLOR.get(cls) if cls else None, fill=PUR, al=CEN, box=True, unlock=True)
    put(ws, f'E{r}', 30 if cls else None, fill=PUR, al=CEN, box=True, unlock=True)
dv = DataValidation(type='list', formula1=COLOR_LIST, allow_blank=True); ws.add_data_validation(dv); dv.add('D11:D26')
dv = DataValidation(type='whole', operator='between', formula1='0', formula2=str(NR), allow_blank=True); ws.add_data_validation(dv); dv.add('E11:E26')
colorcf(ws, 'B11:C26', '$D11')
put(ws, 'A28', '※ 書いた組の数がクラス数になります。「人数」の数だけ番号の枠ができます（最大45人）。', f(**SMALLF))
put(ws, 'A29', '※「名簿での表示」は「6-1」などに書きかえてもかまいません。「色」は組の文字の色です（空欄なら黒）。', f(**SMALLF))
put(ws, 'H9', '委員会の設定（上から順に「委員会①、②…」のシートに入ります）', f(bold=True))
for c, h in zip('HIJK', ['No.', '委員会名', '担当', '活動場所']):
    put(ws, f'{c}10', h, f(bold=True), HDR, CEN, True)
for i in range(NC):
    r = 11 + i
    put(ws, f'H{r}', i + 1, fill=GRY, al=CEN, box=True)
    d = COMM[i] if i < len(COMM) else (None,) * 3
    for c, v in zip('IJK', d):
        put(ws, f'{c}{r}', v, fill=PUR, al=CEN, box=True, unlock=True)
cf(ws, 'I11:I30', 'AND(I11<>"",COUNTIF($I$11:$I$30,I11)>1)', RED)
protect(ws)

# ---------------- 3・4. 6年・5年 ----------------
DX_ON = Border(left=thin, right=thin, top=thin, bottom=thin)
for slot in (1, 2):
    g = (6, 5)[slot - 1]
    ws = sheet(f'{g}年')
    put(ws, 'A1', f'={S}!$B$3&"　"&{S}!$B${6 + slot}&"年　"&{S}!$B$4&"入力"', f(size=16, bold=True))
    ws.merge_cells('E1:J1'); role(ws, 'E1', 'tan')
    for c in range(NCLS):
        sr = 11 + (slot - 1) * 8 + c
        cn, cm, ck = (CL(3 * c + k) for k in (1, 2, 3))
        for w, col in zip((5, 15, 11), (cn, cm, ck)):
            ws.column_dimensions[col].width = w
        ws.merge_cells(f'{cn}3:{ck}3')
        put(ws, f'{cn}3', f'=IF({S}!$B${sr}="","",{S}!$C${sr})', f(bold=True, size=13), HDR, CEN)
        for col in (cn, cm, ck):
            ws[f'{col}3'].border = BOX
        colorcf(ws, f'{cn}3', f'{S}!$D${sr}')
        for col, h in zip((cn, cm, ck), ('番号', '名前', S and '委員会')):
            put(ws, f'{col}4', h, f(bold=True, size=10), HDR, CEN, True)
        rows = DATA.get((g, CLASSES[c])) if c < 4 else None
        for j in range(1, NR + 1):
            r = 4 + j
            put(ws, f'{cn}{r}', f'=IF(AND({S}!$B${sr}<>"",{j}<=N({S}!$E${sr})),{j},"")', f(size=10), al=CEN)
            nm, cm_v = (rows[j - 1] if rows and j <= len(rows) else (None, None))
            put(ws, f'{cm}{r}', nm, f(size=10), al=LFT, unlock=True)
            put(ws, f'{ck}{r}', cm_v, f(size=10), al=CEN, unlock=True)
        rng = f'{cn}5:{ck}{4 + NR}'
        inside = f'AND({S}!${"B"}${sr}<>"",ROW()-4<=N({S}!$E${sr}))'
        cf(ws, f'{ck}5:{ck}{4 + NR}', f'AND({ck}5<>"",COUNTIF({CN},{ck}5)=0)', RED, Font(name=FONT, bold=True, color='C00000'))
        cf(ws, f'{cm}5:{ck}{4 + NR}', f'AND(${cm}5<>"",${ck}5="")', LRED)
        cf(ws, f'{cm}5:{ck}{4 + NR}', inside, YEL, border=DX_ON)
        cf(ws, f'{cn}5:{cn}{4 + NR}', inside, border=DX_ON)
        dv = DataValidation(type='list', formula1=CN, allow_blank=True); ws.add_data_validation(dv); dv.add(f'{ck}5:{ck}{4 + NR}')
    ws.freeze_panes = 'A5'
    ws.print_area = f'A1:X{4 + NR}'
    page(ws, ws.PAPERSIZE_A3, 'landscape', 1)
    protect(ws)

# ---------------- 計算用（6年・5年を1列にまとめる） ----------------
CA, CB = 2, 2 + 2 * NCLS * NR - 1          # 計算用の行範囲
def calc_sheet():
    ws = sheet('計算用')
    for c, h in zip('ABCDEFGHIJ', ['組表示', '番号', '名前', '委員会', 'No', '順', 'キー', '学年', '未入力', '名前なし委員会']):
        put(ws, f'{c}1', h, f(size=9, bold=True))
    r = CA
    for slot in (1, 2):
        sh = f"'{(6, 5)[slot - 1]}年'"
        for c in range(NCLS):
            sr = 11 + (slot - 1) * 8 + c
            cn, cm, ck = (CL(3 * c + k) for k in (1, 2, 3))
            for j in range(1, NR + 1):
                gr = 4 + j
                put(ws, f'A{r}', f'=IF({sh}!{cm}{gr}="","",{S}!$C${sr})', f(size=9))
                put(ws, f'B{r}', f'=IF({sh}!{cm}{gr}="","",{sh}!{cn}{gr})', f(size=9))
                put(ws, f'C{r}', f'=IF({sh}!{cm}{gr}="","",{sh}!{cm}{gr})', f(size=9))
                put(ws, f'D{r}', f'=IF({sh}!{cm}{gr}="","",{sh}!{ck}{gr}&"")', f(size=9))
                put(ws, f'E{r}', f'=IF(D{r}="","",IFERROR(MATCH(D{r},{CN},0),""))', f(size=9))
                put(ws, f'F{r}', f'=IF(E{r}="","",COUNTIF($E${CA}:E{r},E{r}))', f(size=9))
                put(ws, f'G{r}', f'=IF(E{r}="","",E{r}*1000+F{r})', f(size=9))
                put(ws, f'H{r}', f'={S}!$B${6 + slot}', f(size=9))
                put(ws, f'I{r}', f'=IF(AND(C{r}<>"",D{r}=""),1,0)', f(size=9))
                put(ws, f'J{r}', f'=IF(AND(D{r}<>"",E{r}=""),1,0)', f(size=9))
                r += 1
    for k in range(1, NC + 1):
        c = CL(12 + k)
        put(ws, f'{c}1', k, f(size=9, bold=True))
        for j in range(1, NM + 1):
            put(ws, f'{c}{j + 1}', f'=IFERROR(MATCH({k * 1000 + j},$G${CA}:$G${CB},0),"")', f(size=9))
    ws.sheet_state = 'hidden'
    protect(ws)
POS = lambda k, j: f'計算用!${CL(12 + k)}${j + 1}'
def disp(col, p):
    return f'=IF({p}="","",INDEX(計算用!${col}${CA}:${col}${CB},{p}))'
CH, CD = f'計算用!$H${CA}:$H${CB}', f'計算用!$D${CA}:$D${CB}'

# ---------------- 5. 人数チェック ----------------
ws = sheet('人数チェック')
for col, w in zip('ABCDE', [5, 16, 9, 9, 9]):
    ws.column_dimensions[col].width = w
put(ws, 'A1', '人数チェック', f(size=14, bold=True))
ws.merge_cells('D1:H1'); role(ws, 'D1', 'none')
put(ws, 'B3', '委員会が未入力の子', f(bold=True))
put(ws, 'C3', f'={S}!$B$7&"年 "&SUMIFS(計算用!$I${CA}:$I${CB},{CH},{S}!$B$7)&"人"', f(bold=True), al=CEN, box=True)
put(ws, 'D3', f'={S}!$B$8&"年 "&SUMIFS(計算用!$I${CA}:$I${CB},{CH},{S}!$B$8)&"人"', f(bold=True), al=CEN, box=True)
put(ws, 'B4', '設定にない委員会名', f(bold=True))
put(ws, 'C4', f'=SUM(計算用!$J${CA}:$J${CB})&"件"', f(bold=True), al=CEN, box=True)
cf(ws, 'C3:D3', 'VALUE(MID(C3,FIND(" ",C3)+1,LEN(C3)-FIND(" ",C3)-1))>0', RED)
cf(ws, 'C4', 'VALUE(LEFT(C4,LEN(C4)-1))>0', RED)
for c, h in zip('ABCDE', ['No.', '委員会', f'={S}!$B$7&"年"', f'={S}!$B$8&"年"', '合計']):
    put(ws, f'{c}6', h, f(bold=True), HDR, CEN, True)
for i in range(NC):
    r, sr = 7 + i, 11 + i
    put(ws, f'A{r}', i + 1, al=CEN, box=True)
    put(ws, f'B{r}', f'=IF({S}!I{sr}="","",{S}!I{sr})', al=CEN, box=True)
    put(ws, f'C{r}', f'=IF(B{r}="","",COUNTIFS({CH},{S}!$B$7,{CD},B{r}))', al=CEN, box=True)
    put(ws, f'D{r}', f'=IF(B{r}="","",COUNTIFS({CH},{S}!$B$8,{CD},B{r}))', al=CEN, box=True)
    put(ws, f'E{r}', f'=IF(B{r}="","",C{r}+D{r})', f(bold=True), al=CEN, box=True)
put(ws, 'A28', '※ 未入力の子や設定にない委員会名は、6年・5年シートでうすい赤・赤になっています。', f(**SMALLF))
page(ws, ws.PAPERSIZE_A4, 'portrait', 1)
protect(ws)

# ---------------- 6. 一覧名簿 ----------------
ws = sheet('一覧名簿')
T1, T2 = 3, 3 + 4 + NM + 1
for bk in range(10):
    ws.column_dimensions[CL(2 * bk + 1)].width = 5
    ws.column_dimensions[CL(2 * bk + 2)].width = 13
put(ws, 'A1', f'={S}!B3&"　全"&{S}!B4&"担当者および児童名簿"', f(size=16, bold=True))
for (r0, off) in ((T1, 0), (T2, 10)):
    for bk in range(10):
        k = off + bk + 1; sr = 10 + k; qr = 6 + k
        c1, c2 = CL(2 * bk + 1), CL(2 * bk + 2)
        for d in range(3):
            ws.merge_cells(f'{c1}{r0+d}:{c2}{r0+d}')
        put(ws, f'{c1}{r0}', f'=IF({S}!$I${sr}="","",{S}!$I${sr}&"（"&人数チェック!$E${qr}&"名）")', f(bold=True, size=10), HDR, CEN)
        put(ws, f'{c1}{r0+1}', f'=IF({S}!$I${sr}="","","担当："&{S}!$J${sr})', f(size=9), al=CEN)
        put(ws, f'{c1}{r0+2}', f'=IF({S}!$I${sr}="","","場所："&{S}!$K${sr})', f(size=9), al=CEN)
        put(ws, f'{c1}{r0+3}', '組', f(bold=True, size=10), HDR, CEN)
        put(ws, f'{c2}{r0+3}', '名前', f(bold=True, size=10), HDR, CEN)
        for j in range(1, NM + 1):
            r = r0 + 3 + j
            put(ws, f'{c1}{r}', disp('A', POS(k, j)), f(size=9), al=CEN, box=True)
            put(ws, f'{c2}{r}', disp('C', POS(k, j)), f(size=9), al=LFT, box=True)
        for rr in range(r0, r0 + 4):
            for cc in (c1, c2):
                ws[f'{cc}{rr}'].border = BOX
        colorcf(ws, f'{c1}{r0+4}:{c1}{r0+3+NM}', CLR(f'{c1}{r0+4}'))
role(ws, 'V1', 'none'); ws.column_dimensions['V'].width = 30
ws.print_area = f'A1:T{T2 + 3 + NM}'
ws.row_breaks.append(Break(id=T2 - 1))
page(ws, ws.PAPERSIZE_A3, 'landscape', 0)
ws.page_margins.left = ws.page_margins.right = 0.4
ws.page_margins.top = ws.page_margins.bottom = 0.4
protect(ws)

# ---------------- 7〜26. 委員会①〜⑳ ----------------
MARU = '①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳'
for k in range(1, NC + 1):
    ws = sheet(f'委員会{MARU[k-1]}')
    sr, qr = 10 + k, 6 + k
    for col, w in zip('ABCD', [5, 6, 6, 18]):
        ws.column_dimensions[col].width = w
    for i in range(11):
        ws.column_dimensions[CL(5 + i)].width = 4.6
    ws.merge_cells('A1:O1'); ws.merge_cells('A2:O2')
    put(ws, 'A1', f'=IF({S}!$I${sr}="","",{S}!$B$3&"　"&{S}!$I${sr}&{S}!$B$4&"名簿")', f(size=16, bold=True), al=CEN)
    put(ws, 'A2', (f'=IF({S}!$I${sr}="","","担当："&{S}!$J${sr}&"　　活動場所："&{S}!$K${sr}&"　　人数："&人数チェック!$E${qr}&"名（"'
                   f'&{S}!$B$7&"年"&人数チェック!$C${qr}&"・"&{S}!$B$8&"年"&人数チェック!$D${qr}&"）")'), f(size=11), al=CEN)
    for i, h in enumerate(['No.', '組', '番号', '名前'] + list(MARU[:11])):
        put(ws, f'{CL(i+1)}4', h, f(bold=True), HDR, CEN, True)
    for j in range(1, NM + 1):
        r = 4 + j; p = POS(k, j)
        put(ws, f'A{r}', f'=IF({p}="","",{j})', al=CEN, box=True)
        put(ws, f'B{r}', disp('A', p), al=CEN, box=True)
        put(ws, f'C{r}', disp('B', p), al=CEN, box=True)
        put(ws, f'D{r}', disp('C', p), al=LFT, box=True)
        for i in range(11):
            ws[f'{CL(5+i)}{r}'].border = BOX
        ws.row_dimensions[r].height = 15
    colorcf(ws, f'B5:B{4+NM}', CLR('B5'))
    role(ws, 'Q1', 'none'); ws.column_dimensions['Q'].width = 30
    ws.print_area = f'A1:O{4+NM}'
    page(ws, ws.PAPERSIZE_A4, 'portrait', 1)
    protect(ws)

calc_sheet()
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved')
