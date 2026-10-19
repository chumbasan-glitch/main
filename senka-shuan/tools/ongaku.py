# 教育芸術社「指導計画案（主な学習の流れと題材の評価規準例）」Excel（学年ごと）から単元計画を作る
# 使い方: python3 ongaku.py g1.xlsx g2.xlsx ... > ongaku.json
import json, re, sys, warnings
import openpyxl
warnings.filterwarnings("ignore")

def filled(ws):
    # 結合セルは左上の値を全部のセルに入れる
    v = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None: v[(c.row, c.column)] = c.value
    for rg in ws.merged_cells.ranges:
        top = v.get((rg.min_row, rg.min_col))
        for r in range(rg.min_row, rg.max_row + 1):
            for col in range(rg.min_col, rg.max_col + 1): v[(r, col)] = top
    return v

def clean_mat(s):
    s = re.sub(r"[○◇☆♪]|\(共\)|（共）|\(そだてよう\)|（そだてよう）", "", str(s))
    return re.sub(r"\s+", " ", s).strip()

def lesson_nos(x):
    if x is None: return []
    if isinstance(x, (int, float)): return [abs(int(x))]
    return [int(n) for n in re.findall(r"\d+", str(x))]

def unit_from(top, ex):
    t = filled(top)
    name = re.sub(r"^\s*\d+\s*[\.．]\s*", "", str(t.get((1, 2), "")).strip())
    total = int(re.sub(r"\D", "", str(t.get((2, 3), "")).translate(str.maketrans("０１２３４５６７８９", "0123456789"))) or 0)
    aims = []
    for r in range(3, 12):
        a, b = t.get((r, 1)), t.get((r, 2))
        if a and a != "題材の目標" and r > 3: break
        if b: aims.append(str(b).strip())
    e = filled(ex)
    acts, mats = {}, {}
    cur = []
    for r in range(4, ex.max_row + 1):
        ns = lesson_nos(e.get((r, 3)))
        if ns: cur = ns
        if not cur: continue
        m, d = e.get((r, 2)), e.get((r, 4))
        for n in cur:
            if m:
                cm = clean_mat(m)
                if cm and cm not in mats.setdefault(n, []): mats[n].append(cm)
            if d:
                for line in str(d).split("\n"):
                    line = line.strip()
                    if line and line not in acts.setdefault(n, []): acts[n].append(line)
    total = max([total] + list(acts))
    lessons, short = [], []
    for n in range(1, total + 1):
        if not mats.get(n) and mats.get(n - 1): mats[n] = mats[n - 1]   # 教材名がない時間は、前の時間の教材の続き
        ms = "・".join(mats.get(n, []))
        a = [re.sub(r"^●\s*", "", x) for x in acts.get(n, [])]
        lessons.append((("【" + ms + "】\n" if ms else "") + "\n".join("●" + x for x in a)).strip())
        first = re.sub(r"^<[^>]+>", "", a[0]) if a else ""
        short.append((ms + "：" + first) if ms and first else (ms or first))   # 学習内容のない時間は教材名だけ
    return {"name": name, "month": None, "total": total, "aim": "\n".join(aims), "title": "", "lessons": lessons, "short": short}

out = {}
for path in sys.argv[1:]:
    g = int(re.search(r"(\d)", path.split("/")[-1]).group(1))
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets
    units = []
    for i, s in enumerate(ws):
        if "Top" in s.title and i + 1 < len(ws) and "展開" in ws[i + 1].title:
            units.append(unit_from(s, ws[i + 1]))
    out["ongaku-%d" % g] = {"label": "教育芸術社　%d年　音楽（指導計画案）" % g, "grade": g, "subject": "音楽", "units": units}
print(json.dumps(out, ensure_ascii=False))
