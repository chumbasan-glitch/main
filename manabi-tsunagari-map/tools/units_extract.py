# 品川区版：教科書会社の年間指導計画（公開資料）から，単元と学習指導要領の内容の対応を取り出す
# 使い方: python3 tools/units_extract.py <資料のフォルダ>  → midashi/shinagawa/units_pub.tsv
# 資料（教科書会社のファイル）はリポジトリに入れない。取り出すのは単元名と対応だけ。
import os, re, sys, json
import docx, openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = sys.argv[1]
IDS = {i['id'] for i in json.load(open(os.path.join(ROOT, 'data', 'items.json')))['items']}
F = lambda n: os.path.join(SRC, n)
ZEN = str.maketrans('０１２３４５６７８９ＡＢＣＤ（）', '0123456789ABCD()')
KANA = 'アイウエオカキクケコ'
out = []   # (id, 単元名, 巻, 出典)
parts = {}   # (内容, 単元名, ア/イ) → [（ア）などの記号]

def add(iid, unit, book, src):
    assert iid in IDS, (iid, unit)
    unit = re.sub(r'\s+', ' ', unit.replace('／', ' ')).strip()
    unit = re.sub(r'^[○◯]\s*', '', unit)   # 教科書会社の一覧で小単元に付いている「○」は取る
    if (iid, unit) not in {(o[0], o[1]) for o in out}:
        out.append((iid, unit, book.translate(ZEN), src))

def band(subj, g):
    if subj in ('kokugo', 'ongaku', 'zuko', 'seikatsu', 'doutoku'):
        return {1: '12', 2: '12', 3: '34', 4: '34', 5: '56', 6: '56'}[g]
    return str(g)

def grid(t):
    return [[c.text.strip() for c in r.cells] for r in t.rows]

# ---------- 算数（学校図書）：各単元の「学習指導要領との関連」 ----------
for g in range(1, 7):
    ws = openpyxl.load_workbook(F(f'sansu_{g}.xlsx'), data_only=True).worksheets[0]
    rows = [[str(x) for x in r if x not in (None, '')] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if r]
    for k, r in enumerate(rows):
        if not (len(r) == 1 and r[0].startswith('学習指導要領との関連')):
            continue
        j = k - 1
        while j > 0 and not re.match(r'^[０-９0-9]+[　 ]+\S', rows[j][0]):
            j -= 1
        title = rows[j][0]
        if '\n' in title or k - j > 4 or title.endswith('版'):   # 単元の間の小さなページ（考える力をのばそう など）は入れない
            continue
        vol = ''
        for rr in rows[k:k + 40]:
            m = re.search(r'([上下])\s*[／\n]\s*p', ' '.join(rr))
            if m: vol = m.group(1); break
        name = re.sub(r'^[０-９0-9]+[　 ]+', '', title).strip()
        ref = r[0].translate(ZEN)
        ref = ref.split('内容の取')[0].split('[数学的')[0].split('〔数学的')[0]
        for a, n in re.findall(r'([ABCD])\s*\((\d+)\)', ref):
            add(f'sansu-{g}-{a}-{n}', name, f'{g}年{vol}', '学校図書 年間指導計画作成資料')
        # 「A(1)ア…，A(2)…」のように領域を省いて続く書き方
        for a, rest in re.findall(r'([ABCD])\s*\(\d+\)[^ABCD]*?((?:[・，,]\s*\(\d+\)[^ABCD，,]*)+)', ref):
            for n in re.findall(r'\((\d+)\)', rest):
                add(f'sansu-{g}-{a}-{n}', name, f'{g}年{vol}', '学校図書 年間指導計画作成資料')

# ---------- 国語（光村図書）：指導事項配列表（◎＝主に扱う指導事項，☆＝言語活動） ----------
SEC = {'（1）': 'K1', '（2）': 'K2', '（3）': 'K3'}
for g in range(1, 7):
    d = docx.Document(F(f'kokugo_hairetsu_{g}.docx'))
    vols = [p.text.strip().strip('〔〕') for p in d.paragraphs if re.match(r'〔.+巻〕', p.text.strip())]
    ts = d.tables
    for v in range(len(ts) // 2):
        k_t, s_t = grid(ts[2 * v]), grid(ts[2 * v + 1])
        book = vols[v].replace('巻', '') if v < len(vols) else f'{g}年'
        hk = [i for i, r in enumerate(k_t) if any(re.match('^[ア-ケ]　', c) for c in r[1:])][0]
        hs = [i for i, r in enumerate(s_t) if any(re.match('^[ア-ケ]　', c) for c in r)][0]
        kcols = {}
        for c in range(1, len(k_t[0])):
            sec = SEC.get(k_t[hk - 1][c][:3]); let = k_t[hk][c][:1]
            if sec and let in KANA: kcols[c] = f'kokugo-{band("kokugo", g)}-{sec}-{let}'
        scols = {}
        for c in range(len(s_t[0])):
            area = s_t[hs - 2][c][:1].translate(ZEN); kind = s_t[hs - 1][c]; let = s_t[hs][c][:1]
            if area in 'ABC' and let in KANA:
                scols[c] = f'kokugo-{band("kokugo", g)}-{area}-' + ('2' if '言語活動' in kind else let)
        names = [r[0] for r in k_t[hk + 1:]]
        srows = s_t[hs + 1:]
        assert len(names) == len(srows), (g, v, len(names), len(srows))
        for name, kr, sr in zip(names, k_t[hk + 1:], srows):
            if not name: continue
            for c, iid in kcols.items():
                if kr[c] in ('◎', '☆'): add(iid, name, book, '光村図書 指導事項配列表')
            for c, iid in scols.items():
                if sr[c] in ('◎', '☆'): add(iid, name, book, '光村図書 指導事項配列表')

# ---------- 書写（光村図書）：書写の単元は国語の「書写に関する事項」に入れる ----------
SHO = {'12': 'ウ', '34': 'エ', '56': 'エ'}
for g in range(1, 7):
    p = F(f'shosha_{g}.docx')
    if not os.path.exists(p): continue
    for r in grid(docx.Document(p).tables[0]):
        if not re.match(r'^\d$', r[0].translate(ZEN)) or int(r[0].translate(ZEN)) != g: continue
        unit = r[3] or r[4]
        if not unit or unit in ('まとめ',): unit = r[4]
        b = band('kokugo', g)
        add(f'kokugo-{b}-K3-{SHO[b]}', '書写：' + unit, f'{g}年', '光村図書 書写 指導事項と授業時数の配当')

# ---------- 社会・家庭・保健（東京書籍）：「学習指導要領の内容」の列 ----------
def tosho(fn, subj, gfix=None):
    d = docx.Document(F(fn))
    for t in d.tables:
        g = grid(t)
        h = [i for i, r in enumerate(g) if '学習指導要領の内容' in r][0]
        hdr = g[h]; ci = hdr.index('学習指導要領の内容')
        ui = [i for i, x in enumerate(hdr) if x in ('小単元名', '題材名')][0]
        gi = hdr.index('学年') if '学年' in hdr else None
        title = next((p.text for p in d.paragraphs if '第' in p.text and '学年' in p.text), '')
        yield g[h + 1:], ci, ui, gi

for gr in (3, 4, 5, 6):
    for rows, ci, ui, gi in tosho(f'shakai_ryakuan_{gr}.docx', 'shakai'):
        for r in rows:
            unit, ref = r[ui], r[ci].translate(ZEN)
            if not unit or re.search('導入|いかす', unit): continue
            unit = re.sub(r'^\d+[\s　]*', '', unit)
            for n in sorted(set(re.findall(r'\((\d+)\)', ref))):
                add(f'shakai-{gr}-X-{n}', unit, f'{gr}年', '東京書籍 年間指導計画（単元一覧表）')
            # 「(2)ア-(ア)(シ)、イ-(ア)」：本文のどの部分（アの（ア）など）に当たるか
            n = top = None
            for mm in re.finditer(r'\((\d+)\)|([アイ])-|\(([ア-ン])\)', ref):
                if mm.group(1): n, top = mm.group(1), None
                elif mm.group(2): top = mm.group(2)
                elif top and n:
                    key = (f'shakai-{gr}-X-{n}', re.sub(r'\s+', ' ', unit.replace('／', ' ')).strip(), top)
                    parts.setdefault(key, [])
                    if mm.group(3) not in parts[key]: parts[key].append(mm.group(3))

d = docx.Document(F('katei_daizai.docx'))
for ti, t in enumerate(d.tables[:2]):
    gr = 5 + ti; g = grid(t)
    h = [i for i, r in enumerate(g) if '学習指導要領の内容' in r][0]
    hdr = g[h]; ci = hdr.index('学習指導要領の内容'); ui = hdr.index('大題材名')
    for r in g[h + 1:]:
        unit = r[ui].split('\n')[0]; ref = r[ci].translate(str.maketrans('ＡＢＣ', 'ABC'))
        if not unit or 'ガイダンス' in unit: continue
        unit = re.sub(r'^\d+[．.]\s*', '', unit)
        for a, n in re.findall(r'([ABC])\s*\((\d)\)', ref):
            add(f'katei-56-{a}-{n}', unit, f'{gr}年', '東京書籍 年間指導計画（題材一覧表）')

HOKEN = {('3', '1'): 'taiiku-3-G-1', ('4', '2'): 'taiiku-4-G-2', ('5', '1'): 'taiiku-5-G-1',
         ('5', '2'): 'taiiku-5-G-2', ('6', '3'): 'taiiku-6-G-3'}
for fn in ('hoken_34.docx', 'hoken_56.docx'):
    g = grid(docx.Document(F(fn)).tables[0])
    for r in g[1:]:
        if not r[0].endswith('年'): continue
        gr = r[0][0].translate(ZEN); ref = r[4].translate(ZEN)
        unit = re.sub(r'^\(\d+\)', '', r[3].translate(ZEN)).strip('・ ')
        for n in set(re.findall(r'^\(?\((\d)\)', ref)):
            add(HOKEN[(gr, n)], '保健：' + unit, f'{gr}年', '東京書籍 年間指導計画（単元一覧表）')

# ---------- 道徳（東京書籍）：教材名と内容項目 ----------
DNO = {'善悪の判断': '01', '正直': '02', '節度': '03', '個性の伸長': '04', '希望と勇気': '05', '真理の探究': '06',
       '親切': '07', '感謝': '08', '礼儀': '09', '友情': '10', '相互理解': '11', '規則の尊重': '12', '公正': '13',
       '勤労': '14', '家族愛': '15', 'よりよい学校生活': '16', '伝統と文化': '17', '国際理解': '18',
       '生命の尊さ': '19', '自然愛護': '20', '感動': '21', 'よりよく生きる喜び': '22'}
DAREA = {**{k: 'A' for k in ('01', '02', '03', '04', '05', '06')}, **{k: 'B' for k in ('07', '08', '09', '10', '11')},
         **{k: 'C' for k in ('12', '13', '14', '15', '16', '17', '18')}, **{k: 'D' for k in ('19', '20', '21', '22')}}
for gr in range(1, 7):
    g = grid(docx.Document(F(f'doutoku_r_{gr}.docx')).tables[0])
    for r in g:
        cells = list(dict.fromkeys(r))
        k = next((i for i, c in enumerate(cells) if any(c.startswith(w) for w in DNO)), None)
        if k is None or k < 2: continue
        name = cells[k - 2] if re.match(r'^\d+$', cells[k - 1]) else cells[k - 1]
        no = next(v for w, v in DNO.items() if cells[k].startswith(w))
        add(f'doutoku-{band("doutoku", gr)}-{DAREA[no]}-{no}', name, f'{gr}年', '東京書籍 年間指導計画（教材一覧表）')

# ---------- 音楽（教育芸術社）：◎（評価に関わる指導事項）の付いた分野 ----------
for gr in range(1, 7):
    g = grid(docx.Document(F(f'ongaku_{gr}.docx')).tables[0])
    sub = g[2]
    cols = {'歌唱': 'A-1', '器楽': 'A-2', '音楽づくり': 'A-3', '鑑賞': 'B-1'}
    for r in g[5:]:
        unit = r[3].replace('\n', ' ')
        if not re.match(r'^\d+\.', unit): continue
        unit = re.sub(r'^\d+\.\s*', '', unit).replace(' ', '')
        for c, x in enumerate(r):
            if x == '◎' and sub[c] in cols:
                add(f'ongaku-{band("ongaku", gr)}-{cols[sub[c]]}', unit, f'{gr}年', '教育芸術社 年間指導計画作成資料')

# ---------- 図画工作（日本文教出版）：題材別カリキュラムの「分野」 ----------
# 造（造形遊び）・絵・立（立体）・工（工作）はA表現の⑴発想や構想と⑵技能。B鑑賞は「鑑」から始まる（鑑賞が中心の）題材だけ。（ ）付きは関連付けて指導する場合なので入れない
if os.path.exists(F('zuko.xlsx')):
    wb = openpyxl.load_workbook(F('zuko.xlsx'), data_only=True)
    for sh in ('１・２上', '１・２下', '３・４上', '３・４下', '５・６上', '５・６下'):
        b = sh[0].translate(ZEN) + sh[2].translate(ZEN)
        for r in wb[sh].iter_rows(values_only=True):
            if not r[0] or not r[3] or r[2] != sh: continue
            name = str(r[3]).replace('\n', ' ').strip(' ／')
            name = re.sub(r'\s*＊(～[^～]*～)?[^～]*', ' ', name).strip()   # 題材名に続く「＊…」の注は除く
            name = re.sub(r'～\s*～', '～ ～', name)
            main = [x for x in str(r[0]).split('・') if not x.startswith('（')]
            ks = (['A-1', 'A-2'] if any(m in ('造', '絵', '立', '工') for m in main) else []) + (['B-1'] if main and main[0] == '鑑' else [])
            for m in main[:1]:
                for k in ks:
                    add(f'zuko-{b}-{k}', name, sh.translate(ZEN).replace('・', '・'), '日本文教出版 評価規準の参考事例・年間指導計画作成資料')

os.makedirs(os.path.join(ROOT, 'midashi', 'shinagawa'), exist_ok=True)
with open(os.path.join(ROOT, 'midashi', 'shinagawa', 'units_parts.tsv'), 'w') as f:
    f.write('# 単元が本文のどの部分に当たるか（東京書籍 社会の「学習指導要領の内容」の列）。tools/units_extract.py で取り出したもの\n')
    f.write('# 内容\t単元名\tア/イ\t記号\n')
    for (iid, unit, top), ls in parts.items():
        f.write('\t'.join([iid, unit, top, ''.join(ls)]) + '\n')
with open(os.path.join(ROOT, 'midashi', 'shinagawa', 'units_pub.tsv'), 'w') as f:
    f.write('# 教科書会社の年間指導計画（公開資料）から tools/units_extract.py で取り出したもの。手で直さない\n')
    f.write('# 内容\t単元名\t巻\t出典\n')
    for o in out:
        f.write('\t'.join(o) + '\n')
from collections import Counter
print(len(out), Counter(o[0].split('-')[0] for o in out))
