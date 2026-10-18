# 算数の年間指導計画（Excel）を読む：単元名・月・時数・各時間のねらい（●の部分だけ）・学習活動
import openpyxl, re, json, sys
Z = str.maketrans('０１２３４５６７８９', '0123456789')
ws = openpyxl.load_workbook(sys.argv[1], data_only=True).worksheets[0]
val = lambda r, c: (str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else '')
def bullets(text, skip_monster=True):
    # 〈 〉の見出しごとに分け、考え方モンスターの部分は除き、●の行だけを残す（●は残す）
    out, skip = [], False
    for line in text.replace('\r', '').split('\n'):
        l = line.strip()
        if not l: continue
        if l.startswith(('〈', '<')):
            skip = skip_monster and 'モンスター' in l
            continue
        if skip: continue
        if l.startswith('●'): out.append(l)
        elif out: out[-1] += l
    return out
HRS = re.compile(r'^[＜<](\d+)月?.*?(\d+)時間[＞>]')   # 「＜10下旬～…」のように「月」が抜けていても読む
PAGES = re.compile(r'^[＜<][^＞>]*[ぺペ]ージ[＞>]')   # 時数のない単元（「＜2ページ＞」）
onlyA = lambda r: val(r, 1) and not any(val(r, c) for c in (2, 3, 4))
units, cur, gcol = [], None, None
for r in range(1, ws.max_row + 1):
    a = val(r, 1)
    m = HRS.match(a.translate(Z))
    if m or PAGES.match(a.translate(Z)):
        # 「＜○月・○時間＞」の1〜2行上が単元名（2行上にも見出しがあれば、2行上が単元名・すぐ上が副題）
        up1, up2 = val(r - 1, 1), val(r - 2, 1)
        name = up2 if (onlyA(r - 2) and onlyA(r - 1) and not HRS.match(up2.translate(Z))) else up1
        cur = {'name': re.sub(r'^[0-9０-９]+[\s　]+', '', name).strip().rstrip('　 '), 'month': int(m.group(1)) if m else None, 'hours': int(m.group(2)) if m else 0, 'rows': [], 'aim': ''}
        units.append(cur); gcol = None; continue
    if not cur: continue
    if a == '目標': cur['aim'] = (val(r, 2) + '\n' + val(r, 4)).strip(); continue
    if a == '小単元' and val(r, 3) == 'ねらい': gcol = (3, 4); continue
    if a == 'ページ' and val(r, 2) == 'ねらい': gcol = (2, 4); continue
    if gcol:
        g = val(r, gcol[0])
        if not g: continue
        if re.match(r'^[〈<]考え方モンスター', g): continue      # 考え方モンスターの行は数えない
        cur['rows'].append({'goal': bullets(g), 'act': bullets(val(r, gcol[1]), False)})
out = []
for u in [u for u in units if u['hours'] > 0]:
    out.append({'name': u['name'], 'month': u['month'], 'hours': u['hours'], 'rows': len(u['rows']), 'aim': u['aim'],
                'short': ['\n'.join(x['goal']) for x in u['rows']], 'lessons': ['\n'.join(x['act']) for x in u['rows']]})
json.dump(out, open(sys.argv[2], 'w'), ensure_ascii=False, indent=1)
print(len(out), 'units', sum(u['hours'] for u in out), 'hours', '| 行の数の合計', sum(u['rows'] for u in out))
for u in out:
    flag = '' if u['rows'] == u['hours'] else '  ★行の数=%d' % u['rows']
    if flag or '-v' in sys.argv: print('  ', u['month'], u['hours'], u['name'], flag)
