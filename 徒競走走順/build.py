"""徒競走の走順Excelを作るスクリプト。

もとのファイル（リレー組分け_v2_白紙.xlsx）を読み込み、
「4人組」シートを「組分け」シートにして、1組4人か5人で組めるようにし、
名前を直接書きかえられる「手直し」シート、印刷用の「決定版」シート、
レースの順番を決める「走順」シート、A4 1枚に印刷する「印刷用」シートを足して、
白紙版とテストデータ入り版の2つを書き出す。

使い方: python3 build.py
"""
import os
import random
import re
import zipfile

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "リレー組分け_v2_白紙.xlsx")
VERSION = "v8"
OUT_BLANK = os.path.join(HERE, f"徒競走走順_{VERSION}_白紙.xlsx")
OUT_TEST = os.path.join(HERE, f"徒競走走順_{VERSION}_テストデータ入り.xlsx")

PEOPLE_LAST = 186      # 計算用の人の行（2〜186）
HEATS = 50             # 男女それぞれ最大50組まで
HEAT_LAST = HEATS + 1  # 計算用の組の表の最後の行
LANES = 5              # 1組の最大人数（表の列の数）
EDIT_TOP = 6           # 手直しシートの1組目の行
EDIT_LAST = EDIT_TOP + HEATS - 1
RACES = 2 * HEATS      # 走順シートのレースの最大数
RACE_TOP = 6           # 走順シートの1レース目の行
# 手直しシートの列: 名前を入れるところ・人数・内訳・入っていない人・コピー用
EDIT = {
    "男": dict(names=list("BCDEF"), size="M", inner="N", missing="R", copy=["U", "V", "W", "X", "Y"]),
    "女": dict(names=list("GHIJK"), size="O", inner="P", missing="S", copy=["Z", "AA", "AB", "AC", "AD"]),
}
HEAT_SHEET = "組分け"

thin = Side(style="thin")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center")
BOLD = Font(bold=True)

# 計算用シートの「組の表」の列（男子・女子）
#   組, 人数, 赤残り, 白残り, 赤人数, 白人数, 赤累計, 白累計
HEAT_COLS = {
    "男": dict(zip("k s rr wr r w cr cw".split(), "S T U V W X Y Z".split())),
    "女": dict(zip("k s rr wr r w cr cw".split(), "AB AC AD AE AF AG AH AI".split())),
}
# 男女ごとの合計セル（計算用）: 赤の人数, 白の人数, 合計人数, 組数, 決定版の組数,
#   手直しの確認の数, 1組の人数（4か5）
TOTALS = {
    "男": dict(red="$Q$2", white="$Q$3", n="$Q$8", h="$Q$9", fh="$Q$13", err="$Q$15", gs="$Q$32"),
    "女": dict(red="$Q$4", white="$Q$5", n="$Q$10", h="$Q$11", fh="$Q$14", err="$Q$16", gs="$Q$33"),
}
# 計算用シートの「手直しの組ごとの赤・白の人数」の列
EDIT_RW = {"男": ("BG", "BH"), "女": ("BI", "BJ")}


def irregular(size, red, white, gs):
    """調整した組（1組の人数でない、または赤白の差が2人以上）かどうかの式。"""
    return f"OR({size}<>{gs},ABS({red}-{white})>1)"


def setup_input(ws):
    """①データ入力シートの左上に「1組の人数」を選ぶ欄を作る。"""
    from openpyxl.worksheet.datavalidation import DataValidation
    ws["A1"] = None
    ws["B1"] = "1組の人数 男子"
    ws["D1"] = "女子"
    for c in ("B1", "D1"):
        ws[c].font = BOLD
        ws[c].alignment = Alignment(horizontal="right", vertical="center", shrink_to_fit=True)
    dv = DataValidation(type="list", formula1='"4,5"', showErrorMessage=True,
                        errorTitle="1組の人数", error="4か5を選んでください")
    ws.add_data_validation(dv)
    for c in ("C1", "E1"):
        ws[c] = 4
        ws[c].font = Font(bold=True, size=12)
        ws[c].fill = PatternFill("solid", fgColor="FFFFF2CC")
        ws[c].border = BOX
        ws[c].alignment = CENTER
        dv.add(c)
    ws.row_dimensions[1].height = 22


def rng(col, last=HEAT_LAST):
    return f"${col}$2:${col}${last}"


def build_calc(ws):
    """計算用シートに、組の人数・赤白の人数を決める表と、一人一人の組を足す。"""
    # 男女ごとの合計人数と組数
    ws["P8"], ws["Q8"] = "男子人数", "=Q2+Q3"
    ws["P32"], ws["Q32"] = "男子の1組の人数", "=IF(①データ入力!$C$1=5,5,4)"
    ws["P33"], ws["Q33"] = "女子の1組の人数", "=IF(①データ入力!$E$1=5,5,4)"
    ws["P34"], ws["Q34"] = "1組の人数の多い方", "=MAX(Q32,Q33)"
    ws["P9"], ws["Q9"] = "男子組数", "=ROUNDUP(Q8/Q32,0)"
    ws["P10"], ws["Q10"] = "女子人数", "=Q4+Q5"
    ws["P11"], ws["Q11"] = "女子組数", "=ROUNDUP(Q10/Q33,0)"

    heads = ["組", "人数", "赤残り", "白残り", "赤人数", "白人数", "赤累計", "白累計"]
    for g, c in HEAT_COLS.items():
        t = TOTALS[g]
        for i, key in enumerate("k s rr wr r w cr cw".split()):
            ws[f"{c[key]}1"] = f"{g}{heads[i]}"
        for row in range(2, HEAT_LAST + 1):
            k = f"{c['k']}{row}"
            s = f"{c['s']}{row}"
            rr = f"{c['rr']}{row}"
            wr = f"{c['wr']}{row}"
            r = f"{c['r']}{row}"
            ws[k] = row - 1
            # 組の人数: 人数を組数でなるべく均等に分ける（多い組を先、少ない組を後にする）
            # 例: 4人組で18人なら 4,4,4,3,3。5人組で47人なら 5×7組, 4×3組
            ws[s] = (f"=IF(OR({t['h']}=0,{k}>{t['h']}),0,"
                     f"INT({t['n']}/{t['h']})+IF({k}<=MOD({t['n']},{t['h']}),1,0))")
            if row == 2:
                ws[rr] = f"={t['red']}"
                ws[wr] = f"={t['white']}"
            else:
                ws[rr] = f"={t['red']}-{c['cr']}{row - 1}"
                ws[wr] = f"={t['white']}-{c['cw']}{row - 1}"
            # 赤の人数: 基本は半分ずつ。3人・5人の組は残りが多い色を1人多くする。
            # 片方の色が足りなければ、もう片方の色で埋める。
            ws[r] = (f"=IF({s}=0,0,MIN({rr},MAX(IF(MOD({s},2)=0,{s}/2,"
                     f"IF({rr}>={wr},({s}+1)/2,({s}-1)/2)),{s}-{wr})))")
            ws[f"{c['w']}{row}"] = f"={s}-{r}"
            prev_cr = "0" if row == 2 else f"{c['cr']}{row - 1}"
            prev_cw = "0" if row == 2 else f"{c['cw']}{row - 1}"
            ws[f"{c['cr']}{row}"] = f"={prev_cr}+{r}"
            ws[f"{c['cw']}{row}"] = f"={prev_cw}+{c['w']}{row}"

    # 一人一人の組・組の中の番号
    ws["AK1"], ws["AL1"], ws["AM1"], ws["AN1"] = "組", "前の組までの人数", "組内番号", "組キー"
    m, f = HEAT_COLS["男"], HEAT_COLS["女"]
    for row in range(2, PEOPLE_LAST + 1):
        J, K = f"J{row}", f"K{row}"
        AK, AL = f"AK{row}", f"AL{row}"

        def by_kind(tmpl):
            """区分（男赤・男白・女赤・女白）ごとに、使う列を入れかえた式を作る。"""
            return (f'IF({J}="男赤",{tmpl(m["cr"])},IF({J}="男白",{tmpl(m["cw"])},'
                    f'IF({J}="女赤",{tmpl(f["cr"])},{tmpl(f["cw"])})))')

        count_before = lambda col: 'COUNTIF(' + rng(col) + ',"<"&' + K + ')'
        ws[AK] = f'=IF({J}="","",{by_kind(count_before)}+1)'
        ws[AL] = (f'=IF({AK}="","",IF({AK}=1,0,'
                  f'{by_kind(lambda col: f"INDEX({rng(col)},{AK}-1)")}))')
        # 赤は組の前の方、白は赤のあとに並べる
        ws[f"AM{row}"] = (f'=IF({AK}="","",IF(RIGHT({J},1)="赤",{K}-{AL},'
                          f'IF(LEFT({J},1)="男",INDEX({rng(m["r"])},{AK}),'
                          f'INDEX({rng(f["r"])},{AK}))+{K}-{AL}))')
        ws[f"AN{row}"] = f'=IF({AK}="","",F{row}&{AK}&"-"&AM{row})'


def build_heat_sheet(wb, name, index, final):
    """組ごとの表のシートを作る。1組の中に赤も白も最大5人まで並べられる形にする。
    final=False は自動の組（組分け）、True は手直しシートの結果（決定版）。"""
    if name in wb.sheetnames:
        wb.remove(wb[name])
    ws = wb.create_sheet(name, index)

    per_slot = 3                        # 名前・クラス・色
    block = 3 + per_slot * LANES        # 組・人数・内訳 ＋ 5人分
    starts = {"男": 1, "女": block + 2}  # 間に1列あける
    last_row = 3 + HEATS

    for g, c0 in starts.items():
        hc = HEAT_COLS[g]
        t = TOTALS[g]
        e = EDIT[g]
        count = t["fh"] if final else t["h"]
        L = lambda i: openpyxl.utils.get_column_letter(c0 + i)
        ws[f"{L(0)}1"] = "男子" if g == "男" else "女子"
        ws[f"{L(0)}1"].font = BOLD
        # 調整した組があるときの注意
        gs = f"計算用!{t['gs']}"
        if final:
            size = f"手直し!${e['size']}${EDIT_TOP}:${e['size']}${EDIT_LAST}"
            red, white = (f"計算用!{rng(c)}" for c in EDIT_RW[g])
            used = f"(ROW({size})-{EDIT_TOP - 1}<=計算用!{count})"
            ws[f"{L(3)}1"] = (f'=IF(計算用!{t["err"]}>0,"※手直しシートに確認が必要なところがあります",'
                              f'IF(SUMPRODUCT({used}*({size}=0))>0,"※だれもいない組があります（赤い組）",'
                              f'IF(SUMPRODUCT({used}*((({size}<>{gs})+(ABS({red}-{white})>1))>0))>0,'
                              f'"※黄色の組は、人数の都合で調整した組です","")))')
        else:
            size, red, white = (f"計算用!{rng(hc[k])}" for k in ("s", "r", "w"))
            ws[f"{L(3)}1"] = (f'=IF(SUMPRODUCT(({size}>0)*((({size}<>{gs})+(ABS({red}-{white})>1))>0))>0,'
                              f'"※黄色の組は、人数の都合で調整した組です","")')
        ws[f"{L(block - 4)}1"] = f'="1組 "&{gs}&"人"'
        ws[f"{L(block - 4)}1"].font = BOLD
        ws[f"{L(3)}1"].font = Font(bold=True, color="FFFF0000")

        for i, label in enumerate(["組", "人数", "内訳"]):
            ws.merge_cells(f"{L(i)}2:{L(i)}3")
            ws[f"{L(i)}2"] = label
        for j in range(LANES):
            a = 3 + j * per_slot
            ws.merge_cells(f"{L(a)}2:{L(a + per_slot - 1)}2")
            ws[f"{L(a)}2"] = f"{j + 1}人目"
            for k, label in enumerate(["名前", "クラス", "色"]):
                ws[f"{L(a + k)}3"] = label
        for rr in (2, 3):
            for i in range(block):
                cell = ws[f"{L(i)}{rr}"]
                cell.border = BOX
                cell.alignment = CENTER
                cell.font = BOLD

        for row in range(4, last_row + 1):
            A = f"${L(0)}{row}"
            ws[f"{L(0)}{row}"] = f'=IF(ROW()-3<=計算用!{count},ROW()-3,"")'
            er = row - 4 + EDIT_TOP   # 手直しシートの同じ組の行
            if final:
                ws[f"{L(1)}{row}"] = f'=IF({A}="","",手直し!{e["size"]}{er})'
                ws[f"{L(2)}{row}"] = f'=IF({A}="","",手直し!{e["inner"]}{er})'
            else:
                ws[f"{L(1)}{row}"] = f'=IF({A}="","",INDEX(計算用!{rng(hc["s"])},{A}))'
                ws[f"{L(2)}{row}"] = (f'=IF({A}="","","赤"&INDEX(計算用!{rng(hc["r"])},{A})'
                                      f'&"白"&INDEX(計算用!{rng(hc["w"])},{A}))')
            for i in range(3):
                ws[f"{L(i)}{row}"].alignment = CENTER
            for j in range(LANES):
                a = 3 + j * per_slot
                if final:
                    nm = f'手直し!{e["names"][j]}{er}'
                    ws[f"{L(a)}{row}"] = f'=IF({A}="","",{nm}&"")'
                    for k, src in ((1, "M"), (2, "I")):
                        ws[f"{L(a + k)}{row}"] = (
                            f'=IF(OR({A}="",{nm}=""),"",IFERROR(INDEX(計算用!${src}$2:${src}${PEOPLE_LAST},'
                            f'MATCH("{g}"&{nm},計算用!$AO$2:$AO${PEOPLE_LAST},0)),"？"))')
                else:
                    key = f'"{g}"&{A}&"-{j + 1}"'
                    for k, src in enumerate(["D", "M", "I"]):
                        ws[f"{L(a + k)}{row}"] = (
                            f'=IF({A}="","",IFERROR(INDEX(計算用!${src}$2:${src}${PEOPLE_LAST},'
                            f'MATCH({key},計算用!$AN$2:$AN${PEOPLE_LAST},0)),""))')
                ws[f"{L(a + 1)}{row}"].alignment = CENTER
                ws[f"{L(a + 2)}{row}"].alignment = CENTER

        # 見た目
        ws.column_dimensions[L(0)].width = 5
        ws.column_dimensions[L(1)].width = 5
        ws.column_dimensions[L(2)].width = 9
        for j in range(LANES):
            a = 3 + j * per_slot
            ws.column_dimensions[L(a)].width = 13
            ws.column_dimensions[L(a + 1)].width = 5
            ws.column_dimensions[L(a + 2)].width = 4
        area = f"{L(0)}4:{L(block - 1)}{last_row}"
        top = f"${L(0)}4"
        if final:
            # だれもいない組は赤
            ws.conditional_formatting.add(
                f"{L(0)}4:{L(2)}{last_row}",
                FormulaRule(formula=[f'AND({top}<>"",${L(1)}4=0)'],
                            fill=PatternFill(bgColor="FFFF9999", fill_type="solid"),
                            border=BOX, stopIfTrue=True))
        # 調整した組（1組の人数でない・赤白の差が2人以上）は黄色
        if final:
            r_rng, w_rng = (f"計算用!{rng(c)}" for c in EDIT_RW[g])
        else:
            r_rng, w_rng = f"計算用!{rng(hc['r'])}", f"計算用!{rng(hc['w'])}"
        odd = irregular(f"${L(1)}4", f"INDEX({r_rng},{top})", f"INDEX({w_rng},{top})", gs)
        ws.conditional_formatting.add(
            f"{L(0)}4:{L(2)}{last_row}",
            FormulaRule(formula=[f'AND({top}<>"",{odd})'],
                        fill=PatternFill(bgColor="FFFFF2CC", fill_type="solid"),
                        border=BOX, stopIfTrue=False))
        # 赤の人の「色」は薄い赤
        for j in range(LANES):
            col = L(3 + j * per_slot + 2)
            ws.conditional_formatting.add(
                f"{col}4:{col}{last_row}",
                FormulaRule(formula=[f'{col}4="赤"'],
                            fill=PatternFill(bgColor="FFF8CBAD", fill_type="solid")))
        # 組がある行に罫線
        ws.conditional_formatting.add(
            area, FormulaRule(formula=[f'{top}<>""'], border=BOX))

    ws.column_dimensions[openpyxl.utils.get_column_letter(block + 1)].width = 2
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:3"
    return 2 * block + 1  # 印刷する列数


def build_edit_calc(ws):
    """手直しシートの確認に使う計算（計算用シート）。"""
    last = PEOPLE_LAST
    ws["AO1"], ws["AP1"], ws["AQ1"], ws["AR1"], ws["AS1"] = (
        "性別+名前", "手直しでの回数", "入っていない", "男子の入っていない番号", "女子の入っていない番号")
    for row in range(2, last + 1):
        ws[f"AO{row}"] = f'=IF(J{row}="","",F{row}&D{row})'
        ws[f"AP{row}"] = (f'=IF(J{row}="","",IF(F{row}="男",'
                          f'COUNTIF(手直し!${EDIT["男"]["names"][0]}${EDIT_TOP}:${EDIT["男"]["names"][-1]}${EDIT_LAST},D{row}),'
                          f'COUNTIF(手直し!${EDIT["女"]["names"][0]}${EDIT_TOP}:${EDIT["女"]["names"][-1]}${EDIT_LAST},D{row})))')
        ws[f"AQ{row}"] = f'=IF(AND(J{row}<>"",AP{row}=0),1,0)'
        ws[f"AR{row}"] = f'=IF(AND(AQ{row}=1,F{row}="男"),COUNTIFS($AQ$2:AQ{row},1,$F$2:F{row},"男"),"")'
        ws[f"AS{row}"] = f'=IF(AND(AQ{row}=1,F{row}="女"),COUNTIFS($AQ$2:AQ{row},1,$F$2:F{row},"女"),"")'
    keys = f"$AO$2:$AO${last}"
    for g, rows in (("男", (13, 15, 17, 19, 21, 23)), ("女", (14, 16, 18, 20, 22, 24))):
        e = EDIT[g]
        grid = f"手直し!${e['names'][0]}${EDIT_TOP}:${e['names'][-1]}${EDIT_LAST}"
        fh, err, miss, dup, unknown, same = rows
        ws[f"P{fh}"], ws[f"Q{fh}"] = (f"{g}子決定組数",
                                      f'=SUMPRODUCT(MAX(({grid}<>"")*(ROW({grid})-{EDIT_TOP - 1})))')
        ws[f"P{miss}"], ws[f"Q{miss}"] = f"{g}子入っていない", f'=COUNTIFS($AQ$2:$AQ${last},1,$F$2:$F${last},"{g}")'
        ws[f"P{dup}"], ws[f"Q{dup}"] = f"{g}子2回以上", f'=COUNTIFS($AP$2:$AP${last},">1",$F$2:$F${last},"{g}")'
        ws[f"P{unknown}"], ws[f"Q{unknown}"] = (f"{g}子名簿にない名前",
                                                f'=SUMPRODUCT(({grid}<>"")*(COUNTIF({keys},"{g}"&{grid})=0))')
        ws[f"P{same}"], ws[f"Q{same}"] = (f"{g}子同じ名前",
                                          f'=SUMPRODUCT(($F$2:$F${last}="{g}")*({keys}<>"")*(COUNTIF({keys},{keys})>1))')
        ws[f"P{err}"], ws[f"Q{err}"] = f"{g}子確認の数", f"=Q{miss}+Q{dup}+Q{unknown}+Q{same}"
    # 手直しの組ごとの赤・白の人数
    for g, (rc, wc) in EDIT_RW.items():
        e = EDIT[g]
        ws[f"{rc}1"], ws[f"{wc}1"] = f"{g}手直し赤", f"{g}手直し白"
        for row in range(2, HEAT_LAST + 1):
            er = row - 2 + EDIT_TOP
            names = f"手直し!${e['names'][0]}${er}:${e['names'][-1]}${er}"
            for col, team in ((rc, "赤"), (wc, "白")):
                ws[f"{col}{row}"] = (f'=SUMPRODUCT(COUNTIFS($AO$2:$AO${last},"{g}"&{names},'
                                     f'$I$2:$I${last},"{team}"))')


def build_edit_sheet(wb, index):
    """名前を直接書きかえて組を直す「手直し」シート。"""
    ws = wb.create_sheet("手直し", index)
    ws["A1"] = "手直し"
    ws["A1"].font = Font(bold=True, size=14)
    ws["C1"] = (f"① 右の「自動の組（コピー用）」U{EDIT_TOP}:AD{EDIT_LAST} をコピーし、"
                f"B{EDIT_TOP} に「値として貼り付け」します"
                "（右クリック →「貼り付けのオプション」の「値」）。\n"
                "② 名前を切り取り・貼り付けして、組を直します。結果は「決定版」シートに出ます。\n"
                "※名簿やタイムを直したときは、もう一度①からやり直してください。")
    ws["C1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("C1:AD1")
    ws.row_dimensions[1].height = 48
    red_font = Font(bold=True, color="FFFF0000")
    for row2, g in ((2, "男"), (3, "女")):
        t = TOTALS[g]
        rows = {"男": (17, 19, 21, 23), "女": (18, 20, 22, 24)}[g]
        miss, dup, unknown, same = (f"計算用!$Q${r}" for r in rows)
        ws[f"B{row2}"] = f'="{g}子：自動 "&計算用!{t["h"]}&"組 → 手直し "&計算用!{t["fh"]}&"組"'
        ws[f"B{row2}"].font = BOLD
        ws[f"E{row2}"] = (f'=IF(計算用!{t["err"]}=0,IF(計算用!{t["fh"]}=0,"","確認：問題なし"),"確認："'
                          f'&IF({miss}>0,"入っていない人 "&{miss}&"人　","")'
                          f'&IF({dup}>0,"2回以上入っている人 "&{dup}&"人　","")'
                          f'&IF({unknown}>0,"名簿にない名前 "&{unknown}&"こ　","")'
                          f'&IF({same}>0,"同じ名前の人がいます（区別できません）",""))')
        ws[f"E{row2}"].font = red_font

    # 見出し
    ws.merge_cells("A4:A5")
    ws["A4"] = "組"
    heads = []   # 見出しにする列
    for g, e in EDIT.items():
        label = "男子" if g == "男" else "女子"
        for key in ("names", "copy"):
            cols = e[key]
            ws.merge_cells(f"{cols[0]}4:{cols[-1]}4")
            ws[f"{cols[0]}4"] = label
            for j, c in enumerate(cols):
                ws[f"{c}5"] = f"{j + 1}人目"
            heads += cols
        ws.merge_cells(f"{e['size']}4:{e['inner']}4")
        ws[f"{e['size']}4"] = label
        ws[f"{e['size']}5"], ws[f"{e['inner']}5"] = "人数", "内訳"
        ws[f"{e['missing']}5"] = label
        heads += [e["size"], e["inner"], e["missing"]]
    ws.merge_cells(f"{EDIT['男']['missing']}4:{EDIT['女']['missing']}4")
    ws[f"{EDIT['男']['missing']}4"] = "入っていない人"
    copy_cols = EDIT["男"]["copy"] + EDIT["女"]["copy"]
    ws.merge_cells(f"{copy_cols[0]}3:{copy_cols[-1]}3")
    ws[f"{copy_cols[0]}3"] = "自動の組（コピー用）"
    for row in (3, 4, 5):
        for c in ["A"] + heads:
            if row == 3 and c not in copy_cols:
                continue
            cell = ws[f"{c}{row}"]
            cell.font = BOLD
            cell.alignment = CENTER
            cell.border = BOX
    gray = PatternFill("solid", fgColor="FFEDEDED")
    for c in copy_cols:
        for row in (3, 4, 5):
            ws[f"{c}{row}"].fill = gray

    # 自動の組の名前が「組分け」シートのどの列にあるか
    per_slot, block = 3, 3 + 3 * LANES
    auto_col = {g: [openpyxl.utils.get_column_letter(c0 + 3 + j * per_slot) for j in range(LANES)]
                for g, c0 in (("男", 1), ("女", block + 2))}
    for k in range(1, HEATS + 1):
        row = EDIT_TOP + k - 1
        ws[f"A{row}"] = k
        ws[f"A{row}"].alignment = CENTER
        for g, e in EDIT.items():
            names = f"{e['names'][0]}{row}:{e['names'][-1]}{row}"
            ws[f"{e['size']}{row}"] = f'=COUNTIF({names},"?*")'
            rc, wc = EDIT_RW[g]
            cr = row - EDIT_TOP + 2   # 計算用の行
            ws[f"{e['inner']}{row}"] = (f'=IF({e["size"]}{row}=0,"","赤"&計算用!{rc}{cr}'
                                        f'&"白"&計算用!{wc}{cr})')
            ws[f"{e['size']}{row}"].alignment = CENTER
            ws[f"{e['size']}{row}"].number_format = "0;-0;;@"   # 0は表示しない
            ws[f"{e['inner']}{row}"].alignment = CENTER
            ws[f"{e['missing']}{row}"] = (
                f'=IFERROR(INDEX(計算用!$D$2:$D${PEOPLE_LAST},MATCH({k},'
                f'計算用!${"AR" if g == "男" else "AS"}$2:${"AR" if g == "男" else "AS"}${PEOPLE_LAST},0))'
                f'&"（"&INDEX(計算用!$M$2:$M${PEOPLE_LAST},MATCH({k},'
                f'計算用!${"AR" if g == "男" else "AS"}$2:${"AR" if g == "男" else "AS"}${PEOPLE_LAST},0))&"）","")')
            for j, c in enumerate(e["copy"]):
                ws[f"{c}{row}"] = f"={HEAT_SHEET}!{auto_col[g][j]}{row - EDIT_TOP + 4}&\"\""
                ws[f"{c}{row}"].fill = gray
            for c in e["names"]:
                ws[f"{c}{row}"].border = BOX

    # 名前の欄の色: 名簿にない・2回以上は黄色＋赤字、赤の子は薄い赤
    for g, e in EDIT.items():
        first, lastc = e["names"][0], e["names"][-1]
        area = f"{first}{EDIT_TOP}:{lastc}{EDIT_LAST}"
        cell = f"{first}{EDIT_TOP}"
        grid = f"${first}${EDIT_TOP}:${lastc}${EDIT_LAST}"
        key = f'"{g}"&{cell}'
        ws.conditional_formatting.add(area, FormulaRule(
            formula=[f'AND({cell}<>"",OR(COUNTIF(計算用!$AO$2:$AO${PEOPLE_LAST},{key})=0,'
                     f'COUNTIF({grid},{cell})>1))'],
            fill=PatternFill(bgColor="FFFFFF00", fill_type="solid"),
            font=Font(bold=True, color="FFFF0000"), stopIfTrue=True))
        ws.conditional_formatting.add(area, FormulaRule(
            formula=[f'AND({cell}<>"",IFERROR(INDEX(計算用!$I$2:$I${PEOPLE_LAST},'
                     f'MATCH({key},計算用!$AO$2:$AO${PEOPLE_LAST},0))="赤",FALSE))'],
            fill=PatternFill(bgColor="FFF8CBAD", fill_type="solid")))
        # 1組の人数でない・赤白の差が2人以上の組は黄色、0人の組（途中）は赤
        sz = f"${e['size']}{EDIT_TOP}"
        gs = f"計算用!{TOTALS[g]['gs']}"
        area2 = f"{e['size']}{EDIT_TOP}:{e['inner']}{EDIT_LAST}"
        ws.conditional_formatting.add(area2, FormulaRule(
            formula=[f'AND({sz}=0,ROW()-{EDIT_TOP - 1}<計算用!{TOTALS[g]["fh"]})'],
            fill=PatternFill(bgColor="FFFF9999", fill_type="solid"), stopIfTrue=True))
        ws.conditional_formatting.add(area2, FormulaRule(
            formula=[f'AND({sz}>0,' + irregular(
                sz, *(f"INDEX(計算用!{rng(c)},ROW()-{EDIT_TOP - 1})" for c in EDIT_RW[g]), gs) + ")"],
            fill=PatternFill(bgColor="FFFFF2CC", fill_type="solid")))
        # 1組4人のときは、5人目の欄を灰色にする
        c5 = e["names"][4]
        ws.conditional_formatting.add(f"{c5}{EDIT_TOP}:{c5}{EDIT_LAST}", FormulaRule(
            formula=[f'AND({c5}{EDIT_TOP}="",{gs}=4)'],
            fill=PatternFill(bgColor="FFD9D9D9", fill_type="solid")))

    widths = {"A": 5, "L": 2, "M": 5, "N": 8, "O": 5, "P": 8, "Q": 2, "R": 16, "S": 16, "T": 2}
    for c in EDIT["男"]["names"] + EDIT["女"]["names"] + copy_cols:
        widths[c] = 12
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    ws.freeze_panes = f"B{EDIT_TOP}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def build_race_calc(ws):
    """走順の計算（計算用シート）。

    男女それぞれ、組の番号（1が一番速い）を次の順番で走らせる。
      1. 4番目から下の偶数番目を、速い方から（4, 6, 8, …）
      2. 5番目から下の奇数番目を、遅い方から（…, 9, 7, 5）
      3. 最後に 3番目 → 2番目 → 1番目（一番速い組が最終レース）
    例: 10組なら 4, 6, 8, 10, 9, 7, 5, 3, 2, 1。3組以下は遅い組から順に。
    """
    ws["P26"], ws["Q26"] = "先に走る", '=IF(走順!$C$2="女子","女","男")'
    ws["P27"], ws["Q27"] = "後に走る", '=IF(Q26="男","女","男")'
    ws["P28"], ws["Q28"] = "先の組数", '=IF(Q26="男",$Q$13,$Q$14)'
    ws["P29"], ws["Q29"] = "後の組数", '=IF(Q27="男",$Q$13,$Q$14)'
    ws["P30"], ws["Q30"] = "レース数", "=Q28+Q29"
    ws["BA1"], ws["BB1"], ws["BC1"], ws["BD1"], ws["BE1"] = (
        "レース", "男女", "男女ごとのレース", "組数", "組")
    for row in range(2, RACES + 2):
        ws[f"BA{row}"] = row - 1
        ws[f"BB{row}"] = f'=IF(BA{row}>$Q$30,"",IF(BA{row}<=$Q$28,$Q$26,$Q$27))'
        ws[f"BC{row}"] = f'=IF(BB{row}="","",IF(BA{row}<=$Q$28,BA{row},BA{row}-$Q$28))'
        ws[f"BD{row}"] = f'=IF(BB{row}="","",IF(BB{row}="男",$Q$13,$Q$14))'
        p, n = f"BC{row}", f"BD{row}"
        evens = f"INT(({n}-2)/2)"   # 4番目から下の偶数番目の数
        odds = f"INT(({n}-3)/2)"    # 5番目から下の奇数番目の数
        ws[f"BE{row}"] = (f'=IF(BB{row}="","",IF({n}<=3,{n}-{p}+1,'
                          f'IF({p}<={evens},2*{p}+2,'
                          f'IF({p}<={evens}+{odds},({n}-1+MOD({n},2))-2*({p}-{evens}-1),'
                          f'3-({p}-{evens}-{odds}-1)))))')


def build_race_sheet(wb, index):
    """レースの順番の表「走順」。組の中の並び順を、そのまま1〜5コースにする。
    女子の名前は薄ピンクにして、男女を見分けやすくする。"""
    ws = wb.create_sheet("走順", index)
    ws["A1"] = "走順"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A2:B2")
    ws["A2"] = "先に走る"
    ws["A2"].font = BOLD
    ws["A2"].alignment = CENTER
    ws["C2"] = "男子"
    ws["C2"].fill = PatternFill("solid", fgColor="FFFFF2CC")
    ws["C2"].border = BOX
    ws["C2"].alignment = CENTER
    ws["C2"].font = BOLD
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"男子,女子"', showErrorMessage=True,
                        errorTitle="先に走る", error="男子か女子を選んでください")
    ws.add_data_validation(dv)
    dv.add("C2")
    ws["D2"] = "← 男子か女子を選んでください"
    ws.merge_cells("D2:H2")
    ws["E1"] = ('=IF(計算用!$Q$30=0,"","男子 "&計算用!$Q$13&"レース・女子 "&計算用!$Q$14'
                '&"レース　合計 "&計算用!$Q$30&"レース")')
    ws["E1"].font = BOLD
    ws["I2"] = ('=IF(OR(計算用!$Q$15>0,計算用!$Q$16>0),'
                '"※手直しシートに確認が必要なところがあります","")')
    ws["I2"].font = Font(bold=True, color="FFFF0000")

    ws.merge_cells("A4:A5")
    ws["A4"] = "レース"
    for j in range(LANES):
        a = 2 + j * 3
        c1, c3 = openpyxl.utils.get_column_letter(a), openpyxl.utils.get_column_letter(a + 2)
        ws.merge_cells(f"{c1}4:{c3}4")
        ws[f"{c1}4"] = f"{j + 1}コース"
        for k, label in enumerate(["名前", "クラス", "色"]):
            ws.cell(5, a + k, label)
    ncol = 1 + 3 * LANES
    for row in (4, 5):
        for i in range(1, ncol + 1):
            cell = ws.cell(row, i)
            cell.font = BOLD
            cell.border = BOX
            cell.alignment = CENTER

    block = 3 + 3 * LANES
    last_col = openpyxl.utils.get_column_letter(2 * block + 1)
    table = f"決定版!$A$4:${last_col}${3 + HEATS}"   # 決定版の表（男子はA列から、女子は block+1 列右から）
    last = RACE_TOP + RACES - 1
    for row in range(RACE_TOP, last + 1):
        r = row - RACE_TOP + 2   # 計算用の行
        A = f"$A{row}"
        heat = f"計算用!$BE${r}"
        off = f'IF(計算用!$BB${r}="男",0,{block + 1})'
        ws[f"A{row}"] = f'=IF(計算用!$BB${r}="","",計算用!$BA${r})'
        ws[f"A{row}"].alignment = CENTER
        for j in range(LANES):
            for k in range(3):
                col = 2 + j * 3 + k
                ws.cell(row, col).value = f'=IF({A}="","",INDEX({table},{heat},{off}+{4 + j * 3 + k}))'
                if k:
                    ws.cell(row, col).alignment = CENTER

    lastc = openpyxl.utils.get_column_letter(ncol)
    top = f"$A{RACE_TOP}"
    # 女子の名前は薄ピンク、赤の子の「色」は薄い赤
    girl = f'INDEX(計算用!$BB$2:$BB${RACES + 1},ROW()-{RACE_TOP - 1})="女"'
    for j in range(LANES):
        name = openpyxl.utils.get_column_letter(2 + j * 3)
        ws.conditional_formatting.add(f"{name}{RACE_TOP}:{name}{last}", FormulaRule(
            formula=[f'AND({name}{RACE_TOP}<>"",{girl})'],
            fill=PatternFill(bgColor="FFFCE4EC", fill_type="solid"), border=BOX))
        c = openpyxl.utils.get_column_letter(4 + j * 3)
        ws.conditional_formatting.add(f"{c}{RACE_TOP}:{c}{last}", FormulaRule(
            formula=[f'{c}{RACE_TOP}="赤"'],
            fill=PatternFill(bgColor="FFF8CBAD", fill_type="solid"), border=BOX))
    ws.conditional_formatting.add(f"A{RACE_TOP}:{lastc}{last}", FormulaRule(
        formula=[f'{top}<>""'], border=BOX))

    widths = [6] + [14, 5, 4] * LANES
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = f"A{RACE_TOP}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "4:5"
    return ncol


def build_print_sheet(wb, index):
    """A4 1枚に印刷する、名前だけのシンプルな走順表「印刷用」。
    1行目がコース、2行目から「○レース」と名前。女子の名前は薄ピンク。"""
    ws = wb.create_sheet("印刷用", index)
    big = Font(size=14)
    head = Font(bold=True, size=14)
    ws["A1"].border = BOX
    for j in range(LANES):
        c = ws.cell(1, 2 + j)
        c.value = f"{j + 1}コース" if j < 4 else f'=IF(計算用!$Q$34=5,"{j + 1}コース","")'
        c.font = head
        c.alignment = CENTER
        c.border = BOX
    last = 1 + RACES
    for row in range(2, last + 1):
        src = row - 2 + RACE_TOP   # 走順シートの行
        ws[f"A{row}"] = f'=IF(走順!$A${src}="","",走順!$A${src}&"レース")'
        ws[f"A{row}"].font = head
        ws[f"A{row}"].alignment = CENTER
        for j in range(LANES):
            name = openpyxl.utils.get_column_letter(2 + j * 3)
            c = ws.cell(row, 2 + j)
            c.value = f'=IF($A{row}="","",走順!{name}{src}&"")'
            c.font = big
            c.alignment = CENTER
        ws.row_dimensions[row].height = 28
    ws.row_dimensions[1].height = 28
    lastc = openpyxl.utils.get_column_letter(1 + LANES)
    girl = f'INDEX(計算用!$BB$2:$BB${RACES + 1},ROW()-1)="女"'
    ws.conditional_formatting.add(f"B2:{lastc}{last}", FormulaRule(
        formula=[f'AND(B2<>"",{girl})'],
        fill=PatternFill(bgColor="FFFCE4EC", fill_type="solid"), border=BOX))
    ws.conditional_formatting.add(f"A2:{lastc}{last}", FormulaRule(
        formula=['$A2<>""'], border=BOX))
    ws.column_dimensions["A"].width = 11
    for j in range(LANES):
        ws.column_dimensions[openpyxl.utils.get_column_letter(2 + j)].width = 19
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1      # A4 1枚に収める
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.freeze_panes = "A2"


def put_print_areas(path, heat_cols, race_cols):
    """人数に合わせて変わる印刷範囲（openpyxlでは書けないので、あとから書き込む）。"""
    names = (
        f"<definedName name=\"_xlnm.Print_Area\" localSheetId=\"1\">"
        f"OFFSET(組ごとタイム順!$A$1,0,0,3+MAX(計算用!$Q$2:$Q$5),17)</definedName>"
        f"<definedName name=\"_xlnm.Print_Area\" localSheetId=\"2\">"
        f"OFFSET({HEAT_SHEET}!$A$1,0,0,3+MAX(計算用!$Q$9,計算用!$Q$11),{heat_cols})</definedName>"
        f"<definedName name=\"_xlnm.Print_Area\" localSheetId=\"3\">"
        f"手直し!$A$1:${EDIT['女']['missing']}${EDIT_LAST}</definedName>"
        f"<definedName name=\"_xlnm.Print_Area\" localSheetId=\"4\">"
        f"OFFSET(決定版!$A$1,0,0,3+MAX(計算用!$Q$13,計算用!$Q$14),{heat_cols})</definedName>"
        f"<definedName name=\"_xlnm.Print_Area\" localSheetId=\"5\">"
        f"OFFSET(走順!$A$1,0,0,{RACE_TOP - 1}+計算用!$Q$30,1+3*計算用!$Q$34)</definedName>"
        f"<definedName name=\"_xlnm.Print_Area\" localSheetId=\"6\">"
        f"OFFSET(印刷用!$A$1,0,0,1+計算用!$Q$30,1+計算用!$Q$34)</definedName>"
    )
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                xml = data.decode("utf-8")
                xml = re.sub(r'<definedName name="_xlnm.Print_Area"[^>]*>.*?</definedName>', "", xml)
                if "<definedNames>" in xml:
                    xml = xml.replace("<definedNames>", "<definedNames>" + names, 1)
                else:
                    xml = xml.replace("</sheets>", "</sheets><definedNames>" + names + "</definedNames>", 1)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)


LAST = ["佐藤", "鈴木", "高橋", "田中", "伊藤", "渡辺", "山本", "中村", "小林", "加藤",
        "吉田", "山田", "山口", "松本", "井上", "木村", "林", "斎藤", "清水", "森",
        "池田", "橋本", "石川", "前田", "藤田", "岡田", "後藤", "長谷川", "村上", "近藤"]
BOY = ["はると", "そうた", "ゆうと", "れん", "ひなた", "りく", "こうき", "たいが", "そら", "かいと",
       "ゆうま", "あおい", "しょう", "けんた", "だいち", "りょう", "いつき", "さく", "みなと", "つばさ"]
GIRL = ["ゆい", "さくら", "ひまり", "りん", "めい", "あかり", "みお", "ことね", "はな", "えま",
        "ゆな", "こはる", "あおい", "つむぎ", "ほのか", "まお", "しおり", "かのん", "さら", "ひな"]

# テストデータ: クラスごとの（男赤, 男白, 女赤, 女白）の人数
# 男子は赤25・白22（計47人＝4の倍数でない・赤白の差3）
# 女子は赤20・白23（計43人＝4の倍数でない・赤白の差3）
TEST_CLASSES = [(7, 6, 5, 6), (6, 5, 5, 6), (6, 6, 5, 5), (6, 5, 5, 6)]
# テストデータの1組の人数: 男子は4人組、女子は5人組（両方の動きを見られるように）
TEST_GROUP = {"男": 4, "女": 5}


def fill_test(ws):
    """テストデータを入れ、入れた子の一覧（名前, 性別, タイム, 赤白, 通し番号）を返す。"""
    rnd = random.Random(20260929)
    used, kids_all = set(), []
    for ci, (mr, mw, fr, fw) in enumerate(TEST_CLASSES):
        col = 2 + ci * 5
        kids = ([("1", "赤")] * mr + [("1", "白")] * mw + [("2", "赤")] * fr + [("2", "白")] * fw)
        rnd.shuffle(kids)
        for i, (sex, team) in enumerate(kids):
            row = 3 + i
            while True:   # 同姓同名にならないようにする
                name = f"{rnd.choice(LAST)} {rnd.choice(BOY if sex == '1' else GIRL)}"
                if name not in used:
                    used.add(name)
                    break
            base = 9.2 if sex == "1" else 9.6
            time = round(rnd.gauss(base, 0.6), 2)
            ws.cell(row, col, name)
            ws.cell(row, col + 1, int(sex))
            ws.cell(row, col + 2, time)
            ws.cell(row, col + 3, team)
            kids_all.append((name, "男" if sex == "1" else "女", time, team, ci * 37 + i))
    return kids_all


def auto_heats(kids, g, gs):
    """計算用シートと同じ決め方で、組を作る（テストデータの手直し用）。"""
    red = sorted((k for k in kids if k[1] == g and k[3] == "赤"), key=lambda k: (k[2], k[4]))
    white = sorted((k for k in kids if k[1] == g and k[3] == "白"), key=lambda k: (k[2], k[4]))
    n = len(red) + len(white)
    heats, ri, wi = [], 0, 0
    h = -(-n // gs)
    for k in range(1, h + 1):
        s = n // h + (1 if k <= n % h else 0)
        rr, wr = len(red) - ri, len(white) - wi
        t = s // 2 if s % 2 == 0 else ((s + 1) // 2 if rr >= wr else (s - 1) // 2)
        r = min(rr, max(t, s - wr))
        heats.append([x[0] for x in red[ri:ri + r]] + [x[0] for x in white[wi:wi + s - r]])
        ri, wi = ri + r, wi + s - r
    return heats


def fill_test_edit(ws, kids):
    """自動の組を「値として貼り付け」したあと、何人か入れかえた状態にする。"""
    for g, e in EDIT.items():
        heats = auto_heats(kids, g, TEST_GROUP[g])
        if g == "男":
            # 例1: 1組目の2人目と2組目の1人目を入れかえる
            heats[0][1], heats[1][0] = heats[1][0], heats[0][1]
            # 例2: 最後の組（3人組）の1人を、1つ前の組から来た子と入れかえる
            heats[-1][0], heats[-2][3] = heats[-2][3], heats[-1][0]
        for k, names in enumerate(heats):
            for j, name in enumerate(names):
                ws[f"{e['names'][j]}{EDIT_TOP + k}"] = name


def main():
    for out, test in ((OUT_BLANK, False), (OUT_TEST, True)):
        wb = openpyxl.load_workbook(SRC)
        build_calc(wb["計算用"])
        build_edit_calc(wb["計算用"])
        setup_input(wb["①データ入力"])
        idx = wb.sheetnames.index("4人組")
        wb.remove(wb["4人組"])
        heat_cols = build_heat_sheet(wb, HEAT_SHEET, idx, final=False)
        build_edit_sheet(wb, idx + 1)
        build_heat_sheet(wb, "決定版", idx + 2, final=True)
        race_cols = build_race_sheet(wb, idx + 3)
        build_print_sheet(wb, idx + 4)
        build_race_calc(wb["計算用"])
        assert wb.sheetnames == ["①データ入力", "組ごとタイム順", HEAT_SHEET, "手直し",
                                 "決定版", "走順", "印刷用", "計算用"], wb.sheetnames
        if test:
            wb["①データ入力"]["C1"] = TEST_GROUP["男"]
            wb["①データ入力"]["E1"] = TEST_GROUP["女"]
            kids = fill_test(wb["①データ入力"])
            fill_test_edit(wb["手直し"], kids)
        wb.calculation = CalcProperties(fullCalcOnLoad=True)
        wb.active = 0
        wb.save(out)
        put_print_areas(out, heat_cols, race_cols)
        print("wrote", out)


if __name__ == "__main__":
    main()
