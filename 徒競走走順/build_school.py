"""徒競走の走順Excel（本校用・4クラス対抗）を作るスクリプト。

本校のしくみ:
  - 松・竹・梅・月の4クラス対抗。1レースは各クラスから1人ずつの4人。
  - 男女別に、各クラスのタイム順位が同じ子どうしで1レースにする（1位どうし、2位どうし…）。
  - 2年生は男女混合（切替で選ぶ）。
  - 走る順番: 組の番号（1が一番速い）を、偶数番目を速い方から → 奇数番目を遅い方から →
    最後に速い組を遅い順に。男女別は最後の3組、男女混合は最後の5組。
    例: 男女別10組 4,6,8,10,9,7,5,3,2,1 / 男女混合10組 6,8,10,9,7,5,4,3,2,1
  - コースは4つ（始めの番号を選べる）。レースごとにクラスを1コースずつずらす。

シート: データ入力 → 組分け（自動）→ 手直し → 掲示用（男女別のレース名）・最終決定（通しのレース名）
白紙版とテストデータ入り版の2つを書き出す。

使い方: python3 build_school.py
"""
import os
import random
import re
import zipfile

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
VERSION = "v1"
OUT_DIR = os.path.join(HERE, "本校用")
OUT_BLANK = os.path.join(OUT_DIR, f"徒競走走順_本校用_{VERSION}_白紙.xlsx")
OUT_TEST = os.path.join(OUT_DIR, f"徒競走走順_本校用_{VERSION}_テストデータ入り.xlsx")

CLASSES = ["松組", "竹組", "梅組", "月組"]
COLORS = ["FFC6E0B4", "FFBDD7EE", "FFF4B6B6", "FFFFE699"]   # 松=緑, 竹=青, 梅=赤, 月=黄
ROWS = 37                     # 1クラスの最大人数
PEOPLE = 4 * ROWS             # 計算用の人の行は 2〜149
P_LAST = PEOPLE + 1
RACES = 2 * ROWS              # 最大のレース数
R_LAST = RACES + 1
MIX_LABEL = "男女混合（2年生）"

# データ入力シート
IN_TOP = 5                    # 1人目の行
IN_LAST = IN_TOP + ROWS - 1
# 手直しシート
ED_TOP = 6
ED_LAST = ED_TOP + ROWS - 1
GRID = {"A": "BCDE", "B": "GHIJ"}          # 名前を入れるところ（A=男子/男女混合, B=女子）
SIZE = {"A": "L", "B": "M"}
MISSING = {"A": "O", "B": "P"}
COPY = {"A": "RSTU", "B": "VWXY"}
# 組分けシート（A=男子/男女混合, B=女子）。各クラス「名前・タイム」の2列
AUTO_TOP = 5
AUTO_LAST = AUTO_TOP + ROWS - 1
AUTO_START = {"A": 1, "B": 11}

thin = Side(style="thin")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center")
BOLD = Font(bold=True)
RED_BOLD = Font(bold=True, color="FFFF0000")
GRAY = PatternFill("solid", fgColor="FFEDEDED")
INPUT = PatternFill("solid", fgColor="FFFFF2CC")

POSTER_MEMO = ("10月6日の週に\n学年でこの並びを\nやります。\n自分が何レース何コースなのか、\n必ず覚えましょう。\n"
               "また、自分の前後・\n左右の友達も必ず\n覚えましょう。\n１分で並びます。\n\n"
               "授業後コース変更の\n説明をします。\n９月２９日（月）までに\n必要なら修正をして\nおきましょう。\n\n"
               "修正の仕方\n※同じレース内での交換を行う。\n※新しいコースを名前の右側に②のように書き入れる\n"
               "※空欄にも同じレース内ならば、入れてよい。")


def fill(rgb):
    return PatternFill(bgColor=rgb, fill_type="solid")


def box_cells(ws, area, bold=False, center=True):
    for row in ws[area]:
        for c in row:
            c.border = BOX
            if center:
                c.alignment = CENTER
            if bold:
                c.font = BOLD


def grid(block):
    g = GRID[block]
    return f"手直し!${g[0]}${ED_TOP}:${g[-1]}${ED_LAST}"


# ---------------------------------------------------------------- データ入力
def build_input(ws):
    ws["A1"] = "組み方"
    ws["C1"] = "男女別"
    ws["F1"] = "先に走る"
    ws["G1"] = "女子"
    ws["J1"] = "始めのコース"
    ws["K1"] = 2
    for c in ("A1", "F1", "J1"):
        ws[c].font = BOLD
        ws[c].alignment = Alignment(horizontal="right", vertical="center")
    ws.merge_cells("A1:B1")
    for c in ("C1", "G1", "K1"):
        ws[c].fill = INPUT
        ws[c].border = BOX
        ws[c].alignment = CENTER
        ws[c].font = Font(bold=True, size=12)
    ws.merge_cells("C1:D1")
    for cell, values, title in (("C1", f'"男女別,{MIX_LABEL}"', "組み方"),
                                ("G1", '"女子,男子"', "先に走る"),
                                ("K1", '"1,2,3,4,5"', "始めのコース")):
        dv = DataValidation(type="list", formula1=values, showErrorMessage=True, errorTitle=title)
        ws.add_data_validation(dv)
        dv.add(cell)
    ws["N1"] = '=IF(計算用!$U$21>0,"※名前はあるのに、性別かタイムが入っていない人が"&計算用!$U$21&"人います（組に入りません）","")'
    ws["N1"].font = RED_BOLD
    ws.row_dimensions[1].height = 24
    ws["A2"] = ("名簿から名前・性別（1=男子 2=女子）・タイム（秒）を貼り付けてください。"
                "クラス名（3行目）は書きかえられます。男女混合（2年生）のときは性別は空欄でかまいません。")
    ws["A2"].font = Font(color="FF555555")

    ws.merge_cells("A3:A4")
    ws["A3"] = "No."
    sex_dv = DataValidation(type="list", formula1='"1,2"', allow_blank=True)
    ws.add_data_validation(sex_dv)
    for ci, name in enumerate(CLASSES):
        c0 = 2 + ci * 4
        ws.merge_cells(f"{col(c0)}3:{col(c0 + 2)}3")
        ws[f"{col(c0)}3"] = name
        ws[f"{col(c0)}3"].fill = PatternFill("solid", fgColor=COLORS[ci])
        for k, label in enumerate(["名前", "性別", "タイム"]):
            ws.cell(4, c0 + k, label)
        box_cells(ws, f"{col(c0)}3:{col(c0 + 2)}4", bold=True)
        ws.column_dimensions[col(c0)].width = 14
        ws.column_dimensions[col(c0 + 1)].width = 6
        ws.column_dimensions[col(c0 + 2)].width = 8
        ws.column_dimensions[col(c0 + 3)].width = 2
        sex_dv.add(f"{col(c0 + 1)}{IN_TOP}:{col(c0 + 1)}{IN_LAST}")
        for r in range(IN_TOP, IN_LAST + 1):
            for k in range(3):
                cell = ws.cell(r, c0 + k)
                cell.border = BOX
                if k:
                    cell.alignment = CENTER
            ws.cell(r, c0 + 2).number_format = "0.00"
    box_cells(ws, "A3:A4", bold=True)
    for r in range(IN_TOP, IN_LAST + 1):
        ws.cell(r, 1, r - IN_TOP + 1).alignment = CENTER
    ws.column_dimensions["A"].width = 5
    ws.freeze_panes = f"B{IN_TOP}"


# ---------------------------------------------------------------- 計算用
def build_calc(ws):
    heads = {"A": "通し", "B": "クラス", "C": "クラス内行", "D": "名前", "E": "性別(入力)", "F": "性別",
             "G": "タイム", "H": "区分", "I": "有効な区分", "J": "クラス内順位", "K": "自動キー",
             "L": "区分+クラス+名前", "M": "手直しでの回数", "N": "自分のクラスの列での回数",
             "O": "入っていない", "P": "Aの入っていない番号", "Q": "Bの入っていない番号",
             "R": "ちがうクラスの列", "S": "入力不足", "V": "順位(数)", "W": "区分+名前"}
    for c, h in heads.items():
        ws[f"{c}1"] = h
    data = f"データ入力!$B${IN_TOP}:${col(1 + 4 * 4)}${IN_LAST}"
    rng = lambda c: f"${c}$2:${c}${P_LAST}"
    for r in range(2, P_LAST + 1):
        ws[f"A{r}"] = r - 1
        ws[f"B{r}"] = f"=INT((A{r}-1)/{ROWS})+1"
        ws[f"C{r}"] = f"=MOD(A{r}-1,{ROWS})+1"
        ws[f"D{r}"] = f"=TRIM(INDEX({data},C{r},(B{r}-1)*4+1)&\"\")"
        ws[f"E{r}"] = f"=INDEX({data},C{r},(B{r}-1)*4+2)&\"\""
        ws[f"F{r}"] = (f'=IF(OR(E{r}="1",E{r}="１",E{r}="男"),"男",'
                       f'IF(OR(E{r}="2",E{r}="２",E{r}="女"),"女",""))')
        ws[f"G{r}"] = f"=INDEX({data},C{r},(B{r}-1)*4+3)"
        ws[f"H{r}"] = f'=IF(D{r}="","",IF($U$2,"混",F{r}))'
        ws[f"I{r}"] = f'=IF(AND(H{r}<>"",ISNUMBER(G{r})),H{r},"")'
        ws[f"J{r}"] = (f'=IF(I{r}="","",COUNTIFS({rng("I")},I{r},{rng("B")},B{r},{rng("G")},"<"&G{r})'
                       f'+COUNTIFS({rng("I")},I{r},{rng("B")},B{r},{rng("G")},G{r},{rng("A")},"<"&A{r})+1)')
        ws[f"K{r}"] = f'=IF(I{r}="","",I{r}&B{r}&"-"&J{r})'
        ws[f"L{r}"] = f'=IF(I{r}="","",I{r}&B{r}&"|"&D{r})'
        ws[f"M{r}"] = (f'=IF(I{r}="","",IF(I{r}="女",COUNTIF({grid("B")},D{r}),'
                       f'COUNTIF({grid("A")},D{r})))')
        ws[f"N{r}"] = (f'=IF(I{r}="","",IF(I{r}="女",COUNTIF(INDEX({grid("B")},0,B{r}),D{r}),'
                       f'COUNTIF(INDEX({grid("A")},0,B{r}),D{r})))')
        ws[f"O{r}"] = f'=IF(AND(I{r}<>"",M{r}=0),1,0)'
        ws[f"P{r}"] = f'=IF(AND(O{r}=1,I{r}<>"女"),COUNTIFS($O$2:O{r},1,$I$2:I{r},"<>女"),"")'
        ws[f"Q{r}"] = f'=IF(AND(O{r}=1,I{r}="女"),COUNTIFS($O$2:O{r},1,$I$2:I{r},"女"),"")'
        ws[f"R{r}"] = f'=IF(AND(I{r}<>"",M{r}>=1,N{r}=0),1,0)'
        ws[f"S{r}"] = f'=IF(AND(D{r}<>"",I{r}=""),1,0)'
        ws[f"V{r}"] = f'=IF(J{r}="",0,J{r})'
        ws[f"W{r}"] = f'=IF(I{r}="","",I{r}&D{r})'

    gA, gB = grid("A"), grid("B")
    summary = [
        ("男女混合", f'=データ入力!$C$1="{MIX_LABEL}"'),                         # U2
        ("Aの区分", '=IF(U2,"混","男")'),                                        # U3
        ("女子が先", '=AND(NOT(U2),データ入力!$G$1="女子")'),                     # U4
        ("始めのコース", "=IF(ISNUMBER(データ入力!$K$1),データ入力!$K$1,2)"),     # U5
        ("最後に走る速い組の数", "=IF(U2,5,3)"),                                  # U6
        ("Aの組数(手直し)", f'=SUMPRODUCT(MAX(({gA}<>"")*(ROW({gA})-{ED_TOP - 1})))'),   # U7
        ("Bの組数(手直し)", f'=IF(U2,0,SUMPRODUCT(MAX(({gB}<>"")*(ROW({gB})-{ED_TOP - 1}))))'),  # U8
        ("Aの組数(自動)", f"=SUMPRODUCT(MAX(({rng('I')}=$U$3)*{rng('V')}))"),    # U9
        ("Bの組数(自動)", f'=SUMPRODUCT(MAX(({rng("I")}="女")*{rng("V")}))'),    # U10
        ("A入っていない", f'=COUNTIFS({rng("O")},1,{rng("I")},$U$3)'),           # U11
        ("B入っていない", f'=COUNTIFS({rng("O")},1,{rng("I")},"女")'),           # U12
        ("A2回以上", f'=COUNTIFS({rng("M")},">1",{rng("I")},$U$3)'),             # U13
        ("B2回以上", f'=COUNTIFS({rng("M")},">1",{rng("I")},"女")'),             # U14
        ("Aちがうクラス", f'=COUNTIFS({rng("R")},1,{rng("I")},$U$3)'),           # U15
        ("Bちがうクラス", f'=COUNTIFS({rng("R")},1,{rng("I")},"女")'),           # U16
        ("A名簿にない", f'=SUMPRODUCT(({gA}<>"")*(COUNTIF({rng("W")},$U$3&{gA})=0))'),   # U17
        ("B名簿にない", f'=IF(U2,0,SUMPRODUCT(({gB}<>"")*(COUNTIF({rng("W")},"女"&{gB})=0)))'),  # U18
        ("同じ名前", f'=SUMPRODUCT(({rng("W")}<>"")*(COUNTIF({rng("W")},{rng("W")})>1))'),  # U19
        ("確認の数", "=SUM(U11:U18)+U19"),                                         # U20
        ("入力不足", f"=SUM({rng('S')})"),                                         # U21
        ("レース数", "=U7+U8"),                                                    # U22
        ("先に走る組数", "=IF(U4,U8,U7)"),                                          # U23
    ]
    for i, (label, f) in enumerate(summary, start=2):
        ws[f"T{i}"], ws[f"U{i}"] = label, f

    # レースの表（通し番号ごと）
    for c, h in zip("Y Z AA AB AC AD".split(), ["レース(通し)", "表(A/B)", "男女ごとの番号", "組数", "組", "掲示用の名前"]):
        ws[f"{c}1"] = h
    for r in range(2, R_LAST + 1):
        Y, Z, n, N = f"Y{r}", f"Z{r}", f"AA{r}", f"AB{r}"
        ws[Y] = r - 1
        ws[Z] = (f'=IF({Y}>$U$22,"",IF({Y}<=$U$23,IF($U$4,"B","A"),IF($U$4,"A","B")))')
        ws[n] = f'=IF({Z}="","",IF({Y}<=$U$23,{Y},{Y}-$U$23))'
        ws[N] = f'=IF({Z}="","",IF({Z}="A",$U$7,$U$8))'
        T = "$U$6"
        E = f"INT(({N}-{T}+1)/2)"   # T+1番目から下の偶数番目の数
        O = f"INT(({N}-{T})/2)"     # T+2番目から下の奇数番目の数
        ws[f"AC{r}"] = (f'=IF({Z}="","",IF({N}<={T},{N}-{n}+1,IF({n}<={E},{T}-1+2*{n},'
                        f'IF({n}<={E}+{O},({N}-1+MOD({N},2))-2*({n}-{E}-1),{T}-({n}-{E}-{O}-1)))))')
        ws[f"AD{r}"] = f'=IF({Z}="","",IF($U$2,"",IF({Z}="B","女子","男子"))&{n}&"レース")'


def lane_name(r_calc, lane):
    """通し番号の行 r_calc のレースで、左から lane 番目（0〜3）のコースを走る子の名前の式。
    クラスはレースごとに1コースずつずらす: クラス番号 = MOD(lane - n + 1, 4) + 1"""
    n, heat, z = f"計算用!$AA${r_calc}", f"計算用!$AC${r_calc}", f"計算用!$Z${r_calc}"
    cls = f"MOD({lane}-{n}+1,4)+1"
    return (f'IF({z}="","",IF({z}="A",INDEX({grid("A")},{heat},{cls}),'
            f'INDEX({grid("B")},{heat},{cls}))&"")')


def class_color_rules(ws, area, top_row, lane, guard="TRUE"):
    """コース（lane）の列に、そのレースを走るクラスの色をつける。"""
    first = area.split(":")[0]
    n = f"INDEX(計算用!$AA$2:$AA${R_LAST},ROW()-{top_row - 1})"
    for k in range(4):
        ws.conditional_formatting.add(area, FormulaRule(
            formula=[f'AND({first}<>"",{guard},MOD({lane}-{n}+1,4)={k})'],
            fill=fill(COLORS[k]), border=BOX))


# ---------------------------------------------------------------- 組分け（自動）
def build_auto(ws):
    ws["A1"] = "組分け（自動）"
    ws["A1"].font = Font(bold=True, size=14)
    ws["E1"] = "各クラスのタイム順位が同じ子どうしで1レースにします。直すときは「手直し」シートで。"
    for b, c0 in AUTO_START.items():
        label = '=IF(計算用!$U$2,"男女混合","男子")' if b == "A" else '=IF(計算用!$U$2,"（使いません）","女子")'
        ws.merge_cells(f"{col(c0)}2:{col(c0 + 8)}2")
        ws[f"{col(c0)}2"] = label
        ws.merge_cells(f"{col(c0)}3:{col(c0)}4")
        ws[f"{col(c0)}3"] = "順位"
        for ci in range(4):
            a = c0 + 1 + ci * 2
            ws.merge_cells(f"{col(a)}3:{col(a + 1)}3")
            ws[f"{col(a)}3"] = f"=データ入力!${col(2 + ci * 4)}$3"
            ws[f"{col(a)}3"].fill = PatternFill("solid", fgColor=COLORS[ci])
            ws[f"{col(a)}4"], ws[f"{col(a + 1)}4"] = "名前", "タイム"
            ws.column_dimensions[col(a)].width = 13
            ws.column_dimensions[col(a + 1)].width = 6
        box_cells(ws, f"{col(c0)}2:{col(c0 + 8)}4", bold=True)
        ws.column_dimensions[col(c0)].width = 5
        grp = "計算用!$U$3" if b == "A" else '"女"'
        count = "計算用!$U$9" if b == "A" else "計算用!$U$10"
        for r in range(AUTO_TOP, AUTO_LAST + 1):
            k = r - AUTO_TOP + 1
            rank = f"${col(c0)}{r}"
            ws.cell(r, c0, f'=IF({k}<={count},{k},"")').alignment = CENTER
            for ci in range(4):
                a = c0 + 1 + ci * 2
                key = f'{grp}&{ci + 1}&"-"&{rank}'
                for off, src in ((0, "D"), (1, "G")):
                    ws.cell(r, a + off, f'=IF({rank}="","",IFERROR(INDEX(計算用!${src}$2:${src}${P_LAST},'
                                        f'MATCH({key},計算用!$K$2:$K${P_LAST},0)),""))')
                ws.cell(r, a + 1).number_format = "0.00"
                ws.cell(r, a + 1).alignment = CENTER
        ws.conditional_formatting.add(f"{col(c0)}{AUTO_TOP}:{col(c0 + 8)}{AUTO_LAST}",
                                      FormulaRule(formula=[f'${col(c0)}{AUTO_TOP}<>""'], border=BOX))
    ws.column_dimensions[col(AUTO_START["B"] - 1)].width = 2
    ws.freeze_panes = f"A{AUTO_TOP}"
    setup_page(ws, "landscape", fit_height=0, titles="2:4")


# ---------------------------------------------------------------- 手直し
def build_edit(ws):
    ws["A1"] = "手直し"
    ws["A1"].font = Font(bold=True, size=14)
    ws["C1"] = (f"① 右の「自動の組（コピー用）」{COPY['A'][0]}{ED_TOP}:{COPY['B'][-1]}{ED_LAST} をコピーし、"
                f"B{ED_TOP} に「値として貼り付け」します（右クリック →「貼り付けのオプション」の「値」）。\n"
                "② 名前を切り取り・貼り付けして直します。同じ段（行）の4人が1レースになります。\n"
                "※名簿・タイム・組み方を変えたときは、もう一度①からやり直してください。")
    ws["C1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(f"C1:{COPY['B'][-1]}1")
    ws.row_dimensions[1].height = 48
    for row, b in ((2, "A"), (3, "B")):
        if b == "A":
            label = '=IF(計算用!$U$2,"男女混合","男子")'
            miss, dup, wrong, unk, auto, now = "$U$11", "$U$13", "$U$15", "$U$17", "$U$9", "$U$7"
        else:
            label = '=IF(計算用!$U$2,"","女子")'
            miss, dup, wrong, unk, auto, now = "$U$12", "$U$14", "$U$16", "$U$18", "$U$10", "$U$8"
        ws[f"B{row}"] = (f'=IF(AND(計算用!$U$2,"{b}"="B"),"",{label[1:]}&"：自動 "&計算用!{auto}'
                         f'&"組 → 手直し "&計算用!{now}&"組")')
        ws[f"B{row}"].font = BOLD
        c = lambda x: f"計算用!{x}"
        ws[f"F{row}"] = (f'=IF(AND(計算用!$U$2,"{b}"="B"),"",IF({c(miss)}+{c(dup)}+{c(wrong)}+{c(unk)}=0,'
                         f'IF({c(now)}=0,"","確認：問題なし"),"確認："'
                         f'&IF({c(miss)}>0,"入っていない人 "&{c(miss)}&"人　","")'
                         f'&IF({c(dup)}>0,"2回以上入っている人 "&{c(dup)}&"人　","")'
                         f'&IF({c(wrong)}>0,"ちがうクラスの列に入っている人 "&{c(wrong)}&"人　","")'
                         f'&IF({c(unk)}>0,"名簿にない名前 "&{c(unk)}&"こ　","")))')
        ws[f"F{row}"].font = RED_BOLD
    ws[f"{MISSING['A']}2"] = '=IF(計算用!$U$19>0,"※同じ名前の人がいます（区別できません）","")'
    ws[f"{MISSING['A']}2"].font = RED_BOLD

    ws.merge_cells("A4:A5")
    ws["A4"] = "組"
    heads = ["A"]
    for b in "AB":
        label = '=IF(計算用!$U$2,"男女混合","男子")' if b == "A" else '=IF(計算用!$U$2,"（使いません）","女子")'
        for cols in (GRID[b], COPY[b]):
            ws.merge_cells(f"{cols[0]}4:{cols[-1]}4")
            ws[f"{cols[0]}4"] = label
            for ci, c in enumerate(cols):
                ws[f"{c}5"] = f"=データ入力!${col(2 + ci * 4)}$3"
                ws[f"{c}5"].fill = PatternFill("solid", fgColor=COLORS[ci])
            heads += list(cols)
        ws[f"{SIZE[b]}5"] = "人数"
        ws[f"{MISSING[b]}5"] = label
        heads += [SIZE[b], MISSING[b]]
    ws.merge_cells(f"{SIZE['A']}4:{SIZE['B']}4")
    ws[f"{SIZE['A']}4"] = "人数"
    ws[f"{SIZE['A']}5"] = '=IF(計算用!$U$2,"混合","男子")'
    ws[f"{SIZE['B']}5"] = '=IF(計算用!$U$2,"","女子")'
    ws.merge_cells(f"{MISSING['A']}4:{MISSING['B']}4")
    ws[f"{MISSING['A']}4"] = "入っていない人"
    copy_cols = COPY["A"] + COPY["B"]
    ws.merge_cells(f"{copy_cols[0]}3:{copy_cols[-1]}3")
    ws[f"{copy_cols[0]}3"] = "自動の組（コピー用）"
    for row in (3, 4, 5):
        for c in heads:
            if row == 3 and c not in copy_cols:
                continue
            cell = ws[f"{c}{row}"]
            cell.font = BOLD
            cell.alignment = CENTER
            cell.border = BOX
            if c in copy_cols and row != 5:
                cell.fill = GRAY

    for k in range(1, ROWS + 1):
        r = ED_TOP + k - 1
        ws[f"A{r}"] = k
        ws[f"A{r}"].alignment = CENTER
        for b in "AB":
            ws[f"{SIZE[b]}{r}"] = f'=COUNTIF({GRID[b][0]}{r}:{GRID[b][-1]}{r},"?*")'
            ws[f"{SIZE[b]}{r}"].number_format = "0;-0;;@"
            ws[f"{SIZE[b]}{r}"].alignment = CENTER
            idx = "P" if b == "A" else "Q"
            ws[f"{MISSING[b]}{r}"] = (
                f'=IFERROR(INDEX(計算用!$D$2:$D${P_LAST},MATCH({k},計算用!${idx}$2:${idx}${P_LAST},0))'
                f'&"（"&INDEX(データ入力!$B$3:${col(14)}$3,1,(INDEX(計算用!$B$2:$B${P_LAST},'
                f'MATCH({k},計算用!${idx}$2:${idx}${P_LAST},0))-1)*4+1)&"）","")')
            auto_c0 = AUTO_START[b]
            for ci, c in enumerate(COPY[b]):
                ws[f"{c}{r}"] = f"=組分け!{col(auto_c0 + 1 + ci * 2)}{r - ED_TOP + AUTO_TOP}&\"\""
                ws[f"{c}{r}"].fill = GRAY
            for c in GRID[b]:
                ws[f"{c}{r}"].border = BOX

    # 名前の欄: 名簿にない・ちがうクラス・2回以上は黄色＋赤字
    for b in "AB":
        g = GRID[b]
        grp = "計算用!$U$3" if b == "A" else '"女"'
        for ci, c in enumerate(g):
            area = f"{c}{ED_TOP}:{c}{ED_LAST}"
            cell = f"{c}{ED_TOP}"
            ws.conditional_formatting.add(area, FormulaRule(
                formula=[f'AND({cell}<>"",OR(COUNTIF(計算用!$L$2:$L${P_LAST},{grp}&{ci + 1}&"|"&{cell})=0,'
                         f'COUNTIF(${g[0]}${ED_TOP}:${g[-1]}${ED_LAST},{cell})>1))'],
                fill=fill("FFFFFF00"), font=Font(bold=True, color="FFFF0000"), stopIfTrue=True))
        sz = f"${SIZE[b]}{ED_TOP}"
        now = "計算用!$U$7" if b == "A" else "計算用!$U$8"
        area = f"{SIZE[b]}{ED_TOP}:{SIZE[b]}{ED_LAST}"
        ws.conditional_formatting.add(area, FormulaRule(
            formula=[f'AND({sz}=0,ROW()-{ED_TOP - 1}<{now})'], fill=fill("FFFF9999"), stopIfTrue=True))
        ws.conditional_formatting.add(area, FormulaRule(
            formula=[f'AND({sz}>0,{sz}<4)'], fill=fill("FFFFF2CC")))
    # 男女混合のときは女子の欄を灰色に
    ws.conditional_formatting.add(f"{GRID['B'][0]}{ED_TOP}:{GRID['B'][-1]}{ED_LAST}", FormulaRule(
        formula=["計算用!$U$2"], fill=fill("FFD9D9D9")))

    widths = {"A": 5, "F": 2, "K": 2, SIZE["A"]: 6, SIZE["B"]: 6, "N": 2, "Q": 2}
    for c in GRID["A"] + GRID["B"] + COPY["A"] + COPY["B"]:
        widths[c] = 12
    widths[MISSING["A"]] = widths[MISSING["B"]] = 17
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    ws.freeze_panes = f"B{ED_TOP}"
    setup_page(ws, "landscape", fit_height=0)


# ---------------------------------------------------------------- 掲示用（男女別のレース名）
def build_poster(ws):
    for lane in range(4):
        c = ws.cell(1, 2 + lane, f'=(計算用!$U$5+{lane})&"コース"')
        c.font = Font(bold=True, size=14)
        c.alignment = CENTER
        c.border = BOX
    ws["A1"].border = BOX
    for r in range(2, R_LAST + 1):
        ws[f"A{r}"] = f"=計算用!$AD${r}"
        ws[f"A{r}"].font = Font(size=12)
        ws[f"A{r}"].alignment = CENTER
        for lane in range(4):
            c = ws.cell(r, 2 + lane, "=" + lane_name(r, lane))
            c.font = Font(size=14)
            c.alignment = CENTER
        ws.row_dimensions[r].height = 27.75
    ws.conditional_formatting.add(f"A2:E{R_LAST}", FormulaRule(formula=['$A2<>""'], border=BOX))
    ws.merge_cells("G1:I20")
    ws["G1"] = POSTER_MEMO
    ws["G1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["G1"].font = Font(size=12)
    ws["G1"].border = BOX
    ws.column_dimensions["A"].width = 14
    for c in "BCDE":
        ws.column_dimensions[c].width = 17.86
    ws.column_dimensions["F"].width = 2
    for c in "GHI":
        ws.column_dimensions[c].width = 12
    setup_page(ws, "portrait", fit_height=1)


# ---------------------------------------------------------------- 最終決定（通しのレース名・A4 1枚）
def build_final(ws):
    ws["H1"] = "色"
    ws["H1"].font = BOLD
    ws["H1"].alignment = Alignment(horizontal="right")
    ws["I1"] = "色あり"
    ws["I1"].fill = INPUT
    ws["I1"].border = BOX
    ws["I1"].alignment = CENTER
    ws["I1"].font = Font(bold=True, size=12)
    dv = DataValidation(type="list", formula1='"色あり,色なし"', showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add("I1")
    ws["H2"] = "← 印刷の前に選んでください（この欄は印刷されません）"
    head = Font(bold=True, size=14)
    ws["A1"].border = BOX
    for lane in range(4):
        c = ws.cell(1, 2 + lane, f'=(計算用!$U$5+{lane})&"コース"')
        c.font = head
        c.alignment = CENTER
        c.border = BOX
    for r in range(2, R_LAST + 1):
        ws[f"A{r}"] = f'=IF(計算用!$Z${r}="","",計算用!$Y${r}&"レース")'
        ws[f"A{r}"].font = head
        ws[f"A{r}"].alignment = CENTER
        for lane in range(4):
            c = ws.cell(r, 2 + lane, "=" + lane_name(r, lane))
            c.font = Font(size=14)
            c.alignment = CENTER
        ws.row_dimensions[r].height = 28
    ws.row_dimensions[1].height = 28
    girl = f'AND(NOT(計算用!$U$2),INDEX(計算用!$Z$2:$Z${R_LAST},ROW()-1)="B")'
    ws.conditional_formatting.add(f"A2:A{R_LAST}", FormulaRule(
        formula=[f'AND($A2<>"",{girl})'], font=Font(bold=True, size=14, color="FFFF0000"), border=BOX))
    for lane in range(4):
        c = col(2 + lane)
        class_color_rules(ws, f"{c}2:{c}{R_LAST}", 2, lane, guard='$I$1="色あり"')
    ws.conditional_formatting.add(f"A2:E{R_LAST}", FormulaRule(formula=['$A2<>""'], border=BOX))
    ws.column_dimensions["A"].width = 11
    for c in "BCDE":
        ws.column_dimensions[c].width = 19
    ws.column_dimensions["F"].width = 3
    ws.freeze_panes = "A2"
    setup_page(ws, "portrait", fit_height=1)
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def setup_page(ws, orientation, fit_height, titles=None):
    ws.page_setup.orientation = orientation
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_height
    ws.print_options.horizontalCentered = True
    if titles:
        ws.print_title_rows = titles


def put_print_areas(path, sheetnames):
    """人数に合わせて変わる印刷範囲（openpyxlでは書けないので、あとから書き込む）。"""
    areas = {
        "組分け": f"OFFSET(組分け!$A$1,0,0,{AUTO_TOP - 1}+MAX(計算用!$U$9,計算用!$U$10),19)",
        "手直し": f"手直し!$A$1:${MISSING['B']}${ED_LAST}",
        "掲示用": "OFFSET(掲示用!$A$1,0,0,MAX(20,1+計算用!$U$22),9)",
        "最終決定": "OFFSET(最終決定!$A$1,0,0,1+計算用!$U$22,5)",
    }
    names = "".join(
        f'<definedName name="_xlnm.Print_Area" localSheetId="{sheetnames.index(n)}">{a}</definedName>'
        for n, a in areas.items())
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                xml = data.decode("utf-8")
                xml = re.sub(r'<definedName name="_xlnm.Print_Area"[^>]*>.*?</definedName>', "", xml)
                if "<definedNames>" in xml:
                    xml = xml.replace("<definedNames>", "<definedNames>" + names, 1)
                else:
                    xml = xml.replace("</sheets>", "</sheets><definedNames>" + names + "</definedNames>", 1)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)


# ---------------------------------------------------------------- テストデータ
TEST_COUNTS = (19, 16)   # 各クラスの男子・女子の人数


def fill_test(ws):
    """テストデータ（架空の名前）を入れ、子の一覧（名前, クラス番号, 性別, タイム, 通し番号）を返す。"""
    rnd = random.Random(20261005)
    kids = []
    short = ["松", "竹", "梅", "月"]
    for ci in range(4):
        sexes = [1] * TEST_COUNTS[0] + [2] * TEST_COUNTS[1]
        rnd.shuffle(sexes)
        for i, sex in enumerate(sexes):
            r = IN_TOP + i
            name = f"{short[ci]}生{i + 1:02d}"
            time = round(rnd.uniform(8.0, 15.8), 2)
            c0 = 2 + ci * 4
            ws.cell(r, c0, name)
            ws.cell(r, c0 + 1, sex)
            ws.cell(r, c0 + 2, time)
            kids.append((name, ci + 1, "男" if sex == 1 else "女", time, ci * ROWS + i + 1))
    return kids


def auto_grid(kids, grp):
    """計算用シートと同じ決め方で、クラスごとの順位の表を作る（grp: 男・女・混）。"""
    cols = []
    for ci in range(1, 5):
        mine = sorted((k for k in kids if k[1] == ci and (grp == "混" or k[2] == grp)),
                      key=lambda k: (k[3], k[4]))
        cols.append([k[0] for k in mine])
    n = max((len(c) for c in cols), default=0)
    return [[c[i] if i < len(c) else "" for c in cols] for i in range(n)]


def fill_test_edit(ws, kids):
    """自動の組を「値として貼り付け」したあと、何人か直した状態にする。"""
    for b, grp in (("A", "男"), ("B", "女")):
        rows = auto_grid(kids, grp)
        if b == "A":
            # 例: 男子の1組目と2組目の松組の子を入れかえる
            rows[0][0], rows[1][0] = rows[1][0], rows[0][0]
        for k, names in enumerate(rows):
            for ci, name in enumerate(names):
                if name:
                    ws[f"{GRID[b][ci]}{ED_TOP + k}"] = name


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    for out, test in ((OUT_BLANK, False), (OUT_TEST, True)):
        wb = openpyxl.Workbook()
        ws_in = wb.active
        ws_in.title = "データ入力"
        sheets = {}
        for name in ("組分け", "手直し", "掲示用", "最終決定", "計算用"):
            sheets[name] = wb.create_sheet(name)
        build_input(ws_in)
        build_calc(sheets["計算用"])
        build_auto(sheets["組分け"])
        build_edit(sheets["手直し"])
        build_poster(sheets["掲示用"])
        build_final(sheets["最終決定"])
        sheets["計算用"].sheet_state = "hidden"
        if test:
            kids = fill_test(ws_in)
            fill_test_edit(sheets["手直し"], kids)
        wb.calculation = CalcProperties(fullCalcOnLoad=True)
        wb.active = 0
        wb.save(out)
        put_print_areas(out, wb.sheetnames)
        print("wrote", out)


if __name__ == "__main__":
    main()
