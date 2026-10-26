"""35人分のテストデータを入れたファイルを作る： python3 sample.py 元のファイル 出力ファイル"""
import random, sys
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName

random.seed(2026)
src, out = sys.argv[1], sys.argv[2]
wb = load_workbook(src)
ros, inp = wb["名簿"], wb["×入力"]
sei = ["佐藤", "鈴木", "高橋", "田中", "伊藤", "渡辺", "山本", "中村", "小林", "加藤", "吉田", "山田", "佐々木", "山口", "松本",
       "井上", "木村", "林", "斎藤", "清水", "山崎", "森", "池田", "橋本", "阿部", "石川", "山下", "中島", "石井", "小川",
       "前田", "岡田", "長谷川", "藤田", "後藤"]
mei = ["はると", "ゆい", "そうた", "めい", "ひなた", "さくら", "りく", "あおい", "ゆうと", "ひまり", "いつき", "りん", "そら", "ゆな",
       "たいが", "こはる", "れん", "みお", "かいと", "えま", "はやと", "つむぎ", "あさひ", "ここな", "ゆうま", "ほのか", "しょう",
       "あかり", "けんと", "ゆず", "だいち", "みゆ", "こうき", "さら", "りょう"]
special = {
    4: dict(sib="〇", memo="2年1組"), 11: dict(sib="〇", memo="5年3組", stimes=[("11月26日（木）", "15:00〜", "兄弟①の前")]), 19: dict(sib="〇", memo="1年2組・4年1組", stimes=[("11月25日（水）", "14:00〜", "どちらでも"), ("11月25日（水）", "14:30〜", "兄弟②の前")]),
    27: dict(sib="〇", memo="3年1組"),
    2: dict(hope="遅い時刻"), 8: dict(hope="早い時刻"), 15: dict(hope="早い日"), 23: dict(hope="遅い日"),
    30: dict(hope="遅い時刻"), 33: dict(hope="早い時刻"),
    6: dict(gap="必ず空けたい"), 14: dict(gap="出来れば空けたい"), 21: dict(gap="必ず空けたい"), 34: dict(gap="出来れば空けたい"), 25: dict(gap="空けない"),
}
status = {9: "×なし", 17: "×なし", 25: "×なし", 12: "未提出", 31: "未提出", 35: "未入力"}
D, T = 5, 9
for i in range(1, 36):
    r, b = 4 + i, 5 + (i - 1) * 28
    ros[f"A{r}"] = i
    ros[f"B{r}"] = f"{sei[i-1]} {mei[i-1]}"
    sp = special.get(i, {})
    if "sib" in sp: ros[f"E{r}"], ros[f"F{r}"] = sp["sib"], sp["memo"]
    if "hope" in sp: ros[f"C{r}"] = sp["hope"]
    if "gap" in sp: ros[f"D{r}"] = sp["gap"]
    for k, st in enumerate(sp.get("stimes", [])):      # 兄弟①〜③：日付・時刻・順番
        for c, v in zip(range(7 + 3 * k, 10 + 3 * k), st):
            ros.cell(r, c).value = v
    st = status.get(i)
    if st == "×なし": inp[f"F{b}"] = "×なし（どこでもOK）"; continue
    if st == "未提出": inp[f"F{b}"] = "未提出"; continue
    if st == "未入力": continue
    ng = set()
    kind = random.random()
    if kind < 0.3:      # 都合の悪い日がまるごとある
        for d in random.sample(range(1, D + 1), random.randint(1, 3)):
            ng |= {(d, t) for t in range(1, T + 1)}
    elif kind < 0.5:    # 早い時間は無理（仕事）
        cut = random.randint(3, 6)
        ng |= {(d, t) for d in range(1, D + 1) for t in range(1, cut)}
    elif kind < 0.6:    # 遅い時間は無理
        cut = random.randint(5, 8)
        ng |= {(d, t) for d in range(1, D + 1) for t in range(cut, T + 1)}
    for _ in range(random.randint(3, 12)):   # ばらばらの×
        ng.add((random.randint(1, D), random.randint(1, T)))
    if len(ng) >= D * T - 2:
        ng = set(list(ng)[: D * T - 6])
    for d, t in ng:
        inp.cell(b + 2 + t, 2 + d).value = "×"
# まとめて×と、先生の都合の例
base = lambda i: 5 + (i - 1) * 28
inp.cell(base(3) + 1, 2 + 2).value = "×"      # 3番：11/25 全部×
inp.cell(base(7) + 2 + 1, 1).value = "×"      # 7番：13:30 全部×
inp.cell(base(7) + 2 + 2, 1).value = "×"      # 7番：13:45 全部×
inp.cell(base(20) + 1, 2 + 5).value = "×"     # 20番：11/30 全部×
for t in (7, 8, 9):
    wb["設定"].cell(10 + t, 12 + 4).value = "×"  # 先生：11/27 15:00〜 出張
# 読み込み直すと消えてしまう「印刷の範囲（日数・枠数に合わせて変わる）」を付け直す
PS = "印刷（兄弟確認）"
wb[PS].defined_names["_xlnm.Print_Area"] = DefinedName("_xlnm.Print_Area", localSheetId=wb.sheetnames.index(PS),
    attr_text="OFFSET('印刷（兄弟確認）'!$B$3,0,0,3+'設定'!$B$6,1+COUNT('設定'!$B$11:$B$20))")
wb.save(out)
print("saved", out)
