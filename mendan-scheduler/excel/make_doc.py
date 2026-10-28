"""マクロを貼り付けるための Word（手順＋マクロの文章）を作る: python3 make_doc.py 版"""
import sys, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ver = sys.argv[1] if len(sys.argv) > 1 else "3.0"
here = os.path.dirname(os.path.abspath(__file__))
code = open(os.path.join(here, "macro.bas"), encoding="utf-8").read().rstrip("\n").split("\n")
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.8)


def font(run, name="ＭＳ ゴシック", size=10.5, bold=False, color=None):
    run.font.name = name; run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)


def para(text, size=10.5, bold=False, color=None, space=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space)
    font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def img(name, w=15):
    doc.add_picture(os.path.join(here, "docimg", name), width=Cm(w))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


para(f"面談日程ツール Ver.{ver}（マクロ版）　マクロの貼り付け方", size=16, bold=True, color="1F3F7A")
para(f"いっしょに届いた「面談日程ツール_Ver{ver}.xlsx」に、このWordのマクロを貼り付けて、マクロ有効ブック（.xlsm）にします。学校のパソコンで1回だけ行います（10分ほど）。", size=10.5)
para("手順", size=13, bold=True, color="1F3F7A")
para(f"① 「面談日程ツール_Ver{ver}.xlsx」を開いて、キーボードの Alt ＋ F11 を押します。マクロを書く画面（VBA）が開きます。")
para("② 左の一覧の「ThisWorkbook」をダブルクリックします。右に白い画面が出ます。")
img("docimg_s1.png")
para("③ このWordの最後の「ここから」の下から「ここまで」の上までを、マウスでなぞって全部選び、Ctrl ＋ C でコピーします。VBAの白い画面をクリックして Ctrl ＋ V で貼り付け、右上の×で閉じます。")
img("docimg_s2.png")
para("④ Excelの［ファイル］→［名前を付けて保存］で、ファイルの種類を「Excel マクロ有効ブック (*.xlsm)」にして保存します。ふつうの .xlsx で保存するとマクロが消えるので、気をつけてください。")
img("docimg_s3.png", 12)
para("⑤ いったん閉じて、保存した .xlsm を開き直します。黄色い帯が出たら［コンテンツの有効化］を押します。これで使えます。")
para("⑥ ほかの先生には、できた .xlsm を校内で配ります。開いて［コンテンツの有効化］を押せば使えます（①〜④は要りません）。")
img("docimg_s4.png")
para("うまく動かないとき", size=13, bold=True, color="1F3F7A")
for t in ["・青いボタンを押しても何も起きない → マクロが有効になっていません。開き直して［コンテンツの有効化］を押してください。.xlsx のまま保存していないかも確かめてください（ファイル名の最後が .xlsm）。",
          "・赤い帯「マクロの実行がブロックされました」が出る → Excelを閉じて、ファイルを右クリック →［プロパティ］→ 下の「許可する」にチェック →［OK］。",
          "・貼り付けたあと「コンパイル エラー」が出る → 貼り付けがとちゅうで切れているかもしれません。白い画面を Ctrl ＋ A で全部選んで消し、もう一度「ここから」〜「ここまで」を全部コピーして貼り付けてください。",
          "・ThisWorkbook 以外（標準モジュールなど）に貼ると、ボタンやダブルクリックが動きません。かならず ThisWorkbook に貼ってください。",
          "・マクロでした操作は Ctrl ＋ Z で戻せません。「調整・決定」は［ひとつ前に戻す］を使ってください。"]:
    para(t, size=10)
para("マクロでできること", size=13, bold=True, color="1F3F7A")
for t in ["・自動案A：［自動案を作る（A・B両方）］［この案を「調整・決定」へ写す］　自動案B：［この案を「調整・決定」へ写す］",
          "・×入力：マスをダブルクリックで×が付く／消える。「児童を選ぶ」で名前を選ぶとその子の表へ。［次の未入力の子へ］",
          "・調整・決定：名前をクリック→入れられる枠に色。ダブルクリック2回で入れ替え（空いたマスなら移動）。［ひとつ前に戻す］（10回まで）",
          "・設定：時間帯・1枠の長さを変えると、×入力とお手紙の使わない行が自動で合う",
          "・名簿：兄弟が〇の子がいると、兄弟①〜③の列が自動で開く",
          "・兄弟確認用・お手紙：［印刷プレビュー］［PDFで保存］（お手紙は1日の枠が13以上なら夏、12以下なら秋の形）",
          "・開いたとき、Excelの計算方法が「手動」になっていたら「自動」に直す"]:
    para(t, size=10)
doc.add_page_break()
para("▼▼▼ ここから（この下の行から、「ここまで」の上の行までを全部コピー） ▼▼▼", size=11, bold=True, color="B23A3A")
for line in code:
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
    font(p.add_run(line if line else ""), name="ＭＳ ゴシック", size=8)
para("▲▲▲ ここまで ▲▲▲", size=11, bold=True, color="B23A3A")
out = os.path.join(here, f"マクロ（貼り付け用）_Ver{ver}.docx")
doc.save(out)
print("saved", out)
